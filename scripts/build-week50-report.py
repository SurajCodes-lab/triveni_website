"""Expand the editorial report from saved project evidence and render DOCX/HTML."""
from pathlib import Path
import csv, json, re, html
from html.parser import HTMLParser
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'reports/week50'
BASE = 'https://www.trivenicabs.in'

def read_json(p): return json.loads(p.read_text(encoding='utf-8'))
def table(headers, rows):
    return '\n'.join(['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---']*len(headers)) + ' |'] + ['| ' + ' | '.join(str(c).replace('|','/') for c in row) + ' |' for row in rows])
def name(slug): return slug.replace('-to-', ' → ').replace('-', ' ').title()

class Meta(HTMLParser):
    def __init__(self): super().__init__(); self.title=''; self.in_title=False; self.description=''; self.href=''; self.planning=False
    def handle_starttag(self, tag, attrs):
        a=dict(attrs)
        if tag=='title': self.in_title=True
        if tag=='meta' and a.get('name')=='description': self.description=a.get('content','')
        if tag=='section' and a.get('id')=='trip-planning': self.planning=True
        if self.planning and tag=='a': self.href=a.get('href','')
    def handle_endtag(self, tag):
        if tag=='title': self.in_title=False
        if tag=='section': self.planning=False
    def handle_data(self, text):
        if self.in_title: self.title+=text

source=OUT/'week50_seo_report.md'
template=OUT/'week50_seo_report.template.md'
if '{{QUERY_TABLE}}' in source.read_text(encoding='utf-8'):
    template.write_text(source.read_text(encoding='utf-8'),encoding='utf-8')
report=template.read_text(encoding='utf-8')
expanded_narrative=OUT/'report-expanded-narrative.md'
expanded_casebooks=OUT/'report-expanded-casebooks.md'
if expanded_narrative.exists() and expanded_casebooks.exists():
    report=report.replace('## Appendix A. The ten improved public pages', expanded_narrative.read_text(encoding='utf-8')+'\n\n## Appendix A. The ten improved public pages')
    report=report.replace('## Appendix E. Evidence and reporting boundaries',expanded_casebooks.read_text(encoding='utf-8')+'\n\n## Appendix J. Evidence and reporting boundaries')
urls=read_json(OUT/'08-ten-pages/urls.json')
routes={r['slug']:r for r in read_json(OUT/'report-route-evidence.json')}
notes=read_json(OUT/'report-planning-evidence.json')
accounts=read_json(ROOT/'scripts/week50-report-page-accounts.json')
with open(r'D:\Desktop\triveni_next\2.0\week49\Queries.csv',encoding='utf-8-sig') as f:
    queries={r['Top queries']:r for r in csv.DictReader(f)}
qrows=[]; impressions=clicks=0; sections=[]
for i,url in enumerate(urls,1):
    slug=url.rsplit('/',1)[1]; r=routes[slug]; q=queries[slug.replace('-',' ')+' distance']
    imp=int(float(q['Impressions'])); click=int(float(q['Clicks'])); impressions+=imp; clicks+=click
    qrows.append([q['Top queries'],f'{imp:,}',str(click),f'{click/imp*100:.2f}%',q['Position']])
    m=Meta(); m.feed((ROOT/'.next/server/app/distance'/f'{slug}.html').read_text(encoding='utf-8'))
    assert m.href and m.title and m.description
    a=accounts[slug]; n=notes[slug]
    destination=('the existing enquiry on the same guide' if m.href.startswith('#') else f'[the existing booking page]({BASE}{m.href})')
    fares=table(['Recorded journey estimate','Value'],[
        ['Road distance',f"Approximately {r['km']} km"],['Driving time',r['duration']],
        ['Sedan',r['fares']['sedan']],['Ertiga',r['fares']['suv']],['Innova Crysta',r['fares']['innova']],['Tempo Traveller',r['fares']['tempo']]])
    sections.append(f'''### 10.{i} {r['origin']} to {r['destination']}

**Public page:** [{r['origin']} to {r['destination']} distance guide]({url})

#### Why this page was selected

The exact query recorded {imp:,} impressions and {click} clicks, with an average position of {q['Position']} in the saved export. This is selection evidence, not a result of the change or a prediction of additional visits.

{a['problem']}

#### The content added to the page

The new section is headed **“{n[0]}”**. Its three published planning points are:

1. {n[1]}
2. {n[2]}
3. {n[3]}

{a['explanation']}

#### How the addition helps a real enquiry

{a['scenario']}

The onward link reaches {destination}. Its click is recorded as planning-panel navigation, while the eventual contact action remains a separate event. The page's original phone, WhatsApp and enquiry options remain available.

#### The listing and the retained estimates

**Published title:** {m.title}

**Published description:** {m.description}

The description introduces the practical planning value as well as the recorded distance and time. The canonical address remains unchanged. These are the page's published metadata; the report does not claim that a search engine has already displayed this exact wording.

{fares}

#### Preservation and checks

{a['boundary']}

The page retained all prior visible text in order and passed mobile and desktop checks, including its booking destination and a single locally captured navigation event per tested click. Public verification confirmed a successful response, the expected canonical address, planning content and tracking markers. The shared estimate and written-quote explanations are present alongside the route-specific advice.
''')
repairs=read_json(OUT/'04-distance-link-targets.json'); repair_rows=[]
for entry in repairs:
    for item in entry['sources']:
        slug=item['source'].rsplit('/',1)[1]; dest=item['after']
        repair_rows.append([f'[{name(slug)}]({BASE}{item["source"]})','Same-page journey enquiry' if dest.startswith('#') else f'[Existing booking page]({BASE}{dest})'])
before=read_json(OUT/'03-after-city-and-links/link-findings.json')
after={r['target'] for r in read_json(OUT/'06-airport-local/link-findings.json')}
airport=[r for r in before if r['target'] not in after]
assert len(repairs)==47 and len(airport)==12 and len(routes)==206
tokens={
 '{{QUERY_TABLE}}':table(['Exact query','Historical impressions','Clicks','Calculated CTR','Average position'],qrows),
 '{{TOTAL_IMPRESSIONS}}':f'{impressions:,}', '{{TOTAL_CLICKS}}':f'{clicks:,}', '{{TOTAL_CTR}}':f'{clicks/impressions*100:.2f}',
 '{{PAGE_ACCOUNTS}}':'\n'.join(sections),
 '{{URL_REGISTER}}':table(['Page','Public URL'],[[name(u.rsplit('/',1)[1]),f'[{u}]({u})'] for u in urls]),
 '{{REPAIR_REGISTER}}':table(['Distance guide','Repaired action destination'],repair_rows),
 '{{AIRPORT_REGISTER}}':table(['Selected transfer','Enquiry page'],[[name(r['target'].lstrip('/')),f'[{r["sources"][0]}]({BASE}{r["sources"][0]})'] for r in airport]),
 '{{DISTANCE_REGISTER}}':table(['Existing distance route','Approx. km','Separate Innova estimate'],[[f'[{r["origin"]} → {r["destination"]}]({BASE}/distance/{r["slug"]})',r['km'],r['fares']['innova']] for r in sorted(routes.values(),key=lambda r:r['slug'])])
}
for key,val in tokens.items(): report=report.replace(key,val)
coverage=OUT/'report-source-coverage.md'
if coverage.exists(): report+='\n\n'+coverage.read_text(encoding='utf-8')
assert '{{' not in report
assert not re.search(r'\b\d{4}-\d{2}-\d{2}\b|\b(?:January|February|March|April|May|June|July|August|September|October|November|December)\b',report)
assert not re.search(r'F4QJy|isepJr|hU8UsC|BUILD_ID|13\.201\.|/home/ubuntu',report)
source.write_text(report,encoding='utf-8')

# A small Markdown renderer for the deliberately limited report vocabulary.
def inline(s):
    s=html.escape(s)
    s=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'<a href="\2">\1</a>',s)
    s=re.sub(r'\*\*([^*]+)\*\*',r'<strong>\1</strong>',s)
    return s

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.75)
sec.left_margin=sec.right_margin=Inches(.75)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5)
normal.paragraph_format.line_spacing=1.18;normal.paragraph_format.space_after=Pt(7)
for lev,size in [(1,22),(2,15),(3,11.5)]:
    st=doc.styles[f'Heading {lev}'];st.font.name='Georgia';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('1B2A4A');st.paragraph_format.space_before=Pt(16);st.paragraph_format.space_after=Pt(8)
doc.core_properties.title='Triveni Cabs — Week 50 Search Performance Report'
doc.core_properties.author='Triveni Cabs';doc.core_properties.subject='Content expansion, pricing consistency and booking journeys'
header=sec.header.paragraphs[0];header.text='TRIVENI CABS  |  WEEK 50 REPORT';header.style='Caption'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('TRIVENI CABS  •  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def rich(p,s):
    pieces=re.split(r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*)',s)
    for bit in pieces:
        m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',bit)
        if m:
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(m[2],RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');prop=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'0E666B');prop.append(color);run.append(prop)
            t=OxmlElement('w:t');t.text=m[1];run.append(t);link.append(run);p._p.append(link)
        elif bit.startswith('**') and bit.endswith('**'):p.add_run(bit[2:-2]).bold=True
        else:p.add_run(bit)

blocks=[]; lines=report.splitlines();i=0; cover=True
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',c) for c in row):rows.append(row)
            i+=1
        blocks.append('<table><thead><tr>'+''.join('<th>'+inline(c)+'</th>' for c in rows[0])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+inline(c)+'</td>' for c in row)+'</tr>' for row in rows[1:])+'</tbody></table>')
        tab=doc.add_table(rows=0,cols=len(rows[0]));tab.style='Light Shading Accent 1'
        for ri,row in enumerate(rows):
            cells=tab.add_row().cells
            pr=cells[0]._tc.getparent().get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            if ri==0:pr.append(OxmlElement('w:tblHeader'))
            for cell,content in zip(cells,row):
                rich(cell.paragraphs[0],content)
                for run in cell.paragraphs[0].runs:run.font.size=Pt(8.5)
                if ri==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'1B2A4A');cell._tc.get_or_add_tcPr().append(shade)
                    for run in cell.paragraphs[0].runs:run.font.color.rgb=RGBColor(255,255,255);run.bold=True
        doc.add_paragraph();continue
    m=re.match(r'^(#{1,4}) (.*)',line)
    if m:
        level=len(m[1]);text=m[2]
        if level==1:
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(text);r.font.name='Georgia';r.font.size=Pt(30 if text=='Week 50' else 23);r.font.color.rgb=RGBColor.from_string('1B2A4A');r.bold=True
            p.paragraph_format.space_before=Pt(28)
            blocks.append('<h1>'+inline(text)+'</h1>')
        else:
            if cover:
                doc.add_page_break();blocks.append('<div class="cover-end"></div>');cover=False
                contents=[x[3:] for x in lines if x.startswith('## ')]
                doc.add_heading('What this report covers',level=1)
                blocks.append('<section class="contents"><h2>What this report covers</h2>')
                for item in contents:doc.add_paragraph(item);blocks.append('<p>'+inline(item)+'</p>')
                doc.add_page_break();blocks.append('</section>')
            p=doc.add_heading(text,level=level-1)
            if level==2:p.paragraph_format.page_break_before=True
            blocks.append(f'<h{level}>'+inline(text)+f'</h{level}>')
        i+=1;continue
    m=re.match(r'^(\d+)\. (.*)',line)
    if m:
        p=doc.add_paragraph(style='List Number');rich(p,m[2]);blocks.append('<p class="numbered">'+inline(line)+'</p>')
    else:
        p=doc.add_paragraph();rich(p,line)
        if cover:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        blocks.append('<p>'+inline(line)+'</p>')
    i+=1
doc.save(OUT/'week50_seo_report.docx')
css='''@page{size:A4;margin:19mm 17mm}*{box-sizing:border-box}body{font:11pt/1.48 Arial,sans-serif;color:#21242b;margin:0}h1,h2,h3{font-family:Georgia,serif;color:#1b2a4a}h1{text-align:center;font-size:28pt;margin:38px 0}h2{font-size:21pt;border-bottom:2px solid #c9a227;padding-bottom:12px;break-before:page}h3{font-size:16pt;margin:25px 0 12px}h4{font-size:12pt;color:#0e666b;margin:19px 0 7px}h2,h3,h4{break-after:avoid}p{margin:0 0 12px;orphans:3;widows:3}a{color:#0e666b;text-decoration:none;overflow-wrap:anywhere}strong{color:#1b2a4a}table{width:100%;border-collapse:collapse;font-size:9pt;line-height:1.35;margin:16px 0 22px}thead{display:table-header-group}th{background:#1b2a4a;color:white;text-align:left}th,td{padding:8px;border:1px solid #ddd7c7;vertical-align:top}tr:nth-child(even){background:#faf6eb}tr{break-inside:avoid}th:first-child,td:first-child{min-width:22%}.contents{break-after:page}.contents p{margin:7px 0}.cover-end{height:0}.numbered{padding-left:14px}'''
(OUT/'week50_seo_report.html').write_text('<!doctype html><html lang="en"><head><meta charset="utf-8"><title>Triveni Cabs — Week 50 Report</title><style>'+css+'</style></head><body>'+''.join(blocks)+'</body></html>',encoding='utf-8')
stats={'words':len(report.split()),'pageAccounts':len(sections),'distanceCoverage':len(routes),'repairedDistanceTargets':len(repairs),'repairedAirportTargets':len(airport),'historicalCohortImpressions':impressions,'historicalCohortClicks':clicks,'absoluteDatesInBody':False,'releaseIdentifiersInBody':False}
(OUT/'report-validation.json').write_text(json.dumps(stats,indent=2),encoding='utf-8')
print(json.dumps(stats))
