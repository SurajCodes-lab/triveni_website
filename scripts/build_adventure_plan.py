# -*- coding: utf-8 -*-
"""Generate the adventure.trivenicabs.in launch plan as a Word .docx."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------- Brand palette ----------
GOLD = RGBColor(0xE6, 0xB0, 0x00)
DARK = RGBColor(0x1A, 0x1A, 0x2E)
SLATE = RGBColor(0x33, 0x41, 0x55)
GREY = RGBColor(0x6B, 0x72, 0x80)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOLD_HEX = "FACF2D"
DARK_HEX = "1A1A2E"
LIGHT_HEX = "FCF7E6"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
style.font.color.rgb = SLATE

def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        e = OxmlElement(f"w:{tag}"); e.set(qn("w:w"), str(val)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)

def no_space(p):
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)

def heading(text, size=15, color=DARK, space_before=16, space_after=6, bar=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before); p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    if bar:
        b = p.add_run("▌ "); b.font.color.rgb = GOLD; b.font.size = Pt(size); b.bold = True
    r = p.add_run(text); r.font.size = Pt(size); r.bold = True; r.font.color.rgb = color
    return p

def subheading(text, color=TEAL):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = color
    return p

def body(text, size=10.5, after=6, color=SLATE, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = color; r.bold = bold
    return p

def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.1
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True; r.font.color.rgb = DARK; r.font.size = Pt(10.5)
        r2 = p.add_run(text); r2.font.size = Pt(10.5); r2.font.color.rgb = SLATE
    else:
        r = p.add_run(text); r.font.size = Pt(10.5); r.font.color.rgb = SLATE
    return p

# ============================================================ COVER
sp = doc.add_paragraph(); sp.paragraph_format.space_before = Pt(40)
tbl = doc.add_table(rows=1, cols=1); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
c = tbl.rows[0].cells[0]; shade_cell(c, DARK_HEX); set_cell_margins(c, 180, 180, 200, 200); c.width = Inches(6.3)
p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; no_space(p)
r = p.add_run("TRIVENI CABS"); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = GOLD
p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER; no_space(p2)
r = p2.add_run("ADVENTURE  •  RISHIKESH"); r.font.size = Pt(11); r.font.color.rgb = WHITE

sp = doc.add_paragraph(); sp.paragraph_format.space_before = Pt(30)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Launch Plan & Proposal"); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = DARK
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; no_space(p)
r = p.add_run("adventure.trivenicabs.in"); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = TEAL
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(8)
r = p.add_run("Adventure Sports  |  Camping  |  Char Dham & Yatra  |  Bike & Scooty Rentals")
r.font.size = Pt(11.5); r.font.color.rgb = GREY

sp = doc.add_paragraph(); sp.paragraph_format.space_before = Pt(24)
facts = doc.add_table(rows=1, cols=3); facts.alignment = WD_TABLE_ALIGNMENT.CENTER
fdata = [("89", "SEO Pages Planned"), ("30+", "Activities & Tours"), ("3", "New Revenue Lines")]
for i, (big, small) in enumerate(fdata):
    cc = facts.rows[0].cells[i]; shade_cell(cc, LIGHT_HEX); set_cell_margins(cc, 120, 120, 80, 80)
    pp = cc.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; no_space(pp)
    rr = pp.add_run(big); rr.bold = True; rr.font.size = Pt(20); rr.font.color.rgb = GOLD
    p2 = cc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER; no_space(p2)
    rr = p2.add_run(small); rr.font.size = Pt(9); rr.font.color.rgb = SLATE

sp = doc.add_paragraph(); sp.paragraph_format.space_before = Pt(40)
for txt, sz, col, it in [("Prepared for: Triveni Cabs Management", 10.5, SLATE, False),
                          ("Prepared by: Digital & Web Team  •  June 2026", 10.5, GREY, False),
                          ("Status: Proposal v1.0  —  Confidential", 9, GREY, True)]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; no_space(p)
    r = p.add_run(txt); r.font.size = Pt(sz); r.font.color.rgb = col; r.italic = it
doc.add_page_break()

# ============================================================ 1. EXEC SUMMARY
heading("1.  Executive Summary")
body("Triveni Cabs is expanding beyond ground transport into experiential travel with a dedicated "
     "adventure vertical on its own subdomain: adventure.trivenicabs.in. The platform will sell "
     "Rishikesh adventure sports (river rafting, bungee jumping, camping, ziplining and more), "
     "Char Dham & Yatra tours, Haridwar tours, and a brand-new self-drive line — bike & scooty "
     "rentals — directly to travellers.")
body("Rishikesh is India's adventure capital and one of the highest-converting travel-intent markets "
     "online. By keeping the brand on a trivenicabs.in subdomain, we inherit the parent domain's SEO "
     "authority, reuse the same Next.js technology stack, and cross-sell our core cab business "
     "(Delhi–Rishikesh transfers, airport pickups, multi-day tours) to every adventure customer — and "
     "vice versa.")
subheading("What this document delivers")
bullet(" A clear business case and scope for the adventure vertical.", "The opportunity —")
bullet(" A complete, SEO-friendly URL list (89 pages) grouped by category, each mapped to a target keyword.", "The full site map —")
bullet(" The content, SEO, backlink, technical and rollout plan to take the site from zero to ranking.", "The execution plan —")
bullet(" Phased timeline and deliverables for sign-off.", "The roadmap —")

# ============================================================ 2. OPPORTUNITY
heading("2.  The Business Opportunity")
subheading("Why Rishikesh, why now")
bullet(" Rishikesh receives millions of adventure and pilgrimage tourists yearly; rafting, bungee, camping, Char Dham and Haridwar are evergreen, high-search, high-margin products.")
bullet(" Most competitors run thin, single-page sites. A structured, deeply-linked, content-rich site outranks them quickly.")
bullet(" Bike & scooty rentals are an under-served, high-repeat category that pairs naturally with adventure trips and weekend travellers.")
subheading("Our unfair advantage")
bullet(" a subdomain of an established, indexed brand ranks faster than a brand-new domain.", "Parent-domain authority — ")
bullet(" every cab customer is an adventure lead; every adventure / yatra booking needs a Delhi–Rishikesh transfer (which we already sell).", "Cross-sell engine — ")
bullet(" we reuse the existing Next.js + schema + internal-linking system already powering 1,900+ live pages.", "Proven tech & SEO playbook — ")
bullet(" the established Triveni brand, ratings and 24x7 support transfer instantly to the new vertical.", "Trust & reviews — ")

# ============================================================ 3. SCOPE
heading("3.  Scope of Services")
subheading("A. Adventure Sports & Camping (Rishikesh)")
body("River Rafting (9 / 16 / 26 / 36 km stretches) • Bungee Jumping • Giant Swing • Reverse Bungee • "
     "Flying Fox • Zipline over the Ganga • Kayaking • Cliff Jumping • Body Surfing • Superman Shot • "
     "Ganga River Expedition • Camping & Bonfire • Rope Course • Rock Climbing • Rappelling • Wall / Tree "
     "Climbing • Sky Cycling • Go Karting • Paintball • Water Zorbing • Mountain Biking • Waterfall "
     "Trekking • Hot-Air Balloon — plus money-making combo packages.", after=4)
subheading("B. Char Dham, Yatra & Haridwar Tours")
body("Char Dham & Do Dham Yatra, Kedarnath, Badrinath, Gangotri, Yamunotri (by cab and by bike), plus "
     "Haridwar tours and the Ganga Aarti experience — a strong pilgrimage SEO and revenue play "
     "originating from Rishikesh.", after=4)
subheading("C. Bike & Scooty Rentals (NEW)")
body("Self-drive two-wheeler rentals: Royal Enfield (Classic 350, Himalayan), KTM Duke, scooters "
     "(Honda Activa, TVS Jupiter), ATV/quad biking, plus guided Himalayan bike-tour routes.", after=4)
subheading("D. Cross-sell with the core cab business")
body("Bundle pages combining Delhi–Rishikesh cab transfers with adventure and yatra packages, "
     "funneling traffic between adventure.trivenicabs.in and the main trivenicabs.in site.", after=4)
doc.add_page_break()

# ============================================================ 4. URL LIST
heading("4.  Website Architecture & Complete SEO-Friendly URL List")
body("Every URL below follows SEO best practice: lowercase, hyphen-separated, keyword-rich, shallow, "
     "human-readable, and free of parameters or IDs. The base domain is "
     "https://adventure.trivenicabs.in. Each page targets a specific search intent.", after=8)

def url_table(rows):
    t = doc.add_table(rows=1, cols=2); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["URL  (path on adventure.trivenicabs.in)", "Primary Target Keyword"]
    widths = [Inches(3.9), Inches(2.5)]
    for i, h in enumerate(headers):
        hc = t.rows[0].cells[i]; shade_cell(hc, DARK_HEX); set_cell_margins(hc); hc.width = widths[i]
        pp = hc.paragraphs[0]; no_space(pp)
        rr = pp.add_run(h); rr.bold = True; rr.font.size = Pt(9.5); rr.font.color.rgb = WHITE
    for idx, (path, kw) in enumerate(rows):
        cells = t.add_row().cells
        if idx % 2 == 1:
            shade_cell(cells[0], "F4F6F8"); shade_cell(cells[1], "F4F6F8")
        for ci in range(2):
            set_cell_margins(cells[ci]); cells[ci].width = widths[ci]
        p0 = cells[0].paragraphs[0]; no_space(p0)
        r0 = p0.add_run(path); r0.font.size = Pt(9); r0.font.name = "Consolas"; r0.font.color.rgb = TEAL
        p1 = cells[1].paragraphs[0]; no_space(p1)
        r1 = p1.add_run(kw); r1.font.size = Pt(9); r1.font.color.rgb = SLATE
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

subheading("4.1  Core & Conversion Pages  (6)")
url_table([
    ("/", "rishikesh adventure sports & rentals"),
    ("/book-now", "book rishikesh adventure online"),
    ("/contact-us", "rishikesh adventure booking contact"),
    ("/about-us", "about triveni adventure rishikesh"),
    ("/reviews", "rishikesh adventure reviews"),
    ("/faqs", "rishikesh adventure faqs"),
])

subheading("4.2  Adventure Activity Hub Pages  (26)")
url_table([
    ("/river-rafting-in-rishikesh", "river rafting in rishikesh"),
    ("/bungee-jumping-in-rishikesh", "bungee jumping in rishikesh"),
    ("/camping-in-rishikesh", "camping in rishikesh"),
    ("/kayaking-in-rishikesh", "kayaking in rishikesh"),
    ("/zipline-in-rishikesh", "zipline rishikesh ganga"),
    ("/giant-swing-in-rishikesh", "giant swing rishikesh"),
    ("/reverse-bungee-in-rishikesh", "reverse bungee rishikesh"),
    ("/flying-fox-in-rishikesh", "flying fox rishikesh"),
    ("/cliff-jumping-in-rishikesh", "cliff jumping rishikesh"),
    ("/body-surfing-in-rishikesh", "body surfing rishikesh"),
    ("/superman-shot-in-rishikesh", "superman zipline rishikesh"),
    ("/scad-jump-in-rishikesh", "scad jump rishikesh"),
    ("/rope-course-in-rishikesh", "rope course rishikesh"),
    ("/rock-climbing-in-rishikesh", "rock climbing rishikesh"),
    ("/rappelling-in-rishikesh", "rappelling in rishikesh"),
    ("/wall-climbing-in-rishikesh", "wall climbing rishikesh"),
    ("/tree-climbing-in-rishikesh", "tree climbing rishikesh"),
    ("/sky-cycling-in-rishikesh", "sky cycling rishikesh"),
    ("/go-karting-in-rishikesh", "go karting rishikesh"),
    ("/paintball-in-rishikesh", "paintball rishikesh"),
    ("/water-zorbing-in-rishikesh", "water zorbing rishikesh"),
    ("/mountain-biking-in-rishikesh", "mountain biking rishikesh"),
    ("/waterfall-trekking-in-rishikesh", "neer waterfall trek rishikesh"),
    ("/ganga-river-expedition-rishikesh", "ganga river expedition rafting"),
    ("/night-camping-bonfire-in-rishikesh", "night camping bonfire rishikesh"),
    ("/hot-air-balloon-in-rishikesh", "hot air balloon rishikesh"),
])

subheading("4.3  River Rafting — Stretch Sub-Pages  (4)")
url_table([
    ("/river-rafting-in-rishikesh/9-km-brahmapuri-to-nim-beach", "9 km rafting rishikesh"),
    ("/river-rafting-in-rishikesh/16-km-shivpuri-to-nim-beach", "16 km rafting rishikesh"),
    ("/river-rafting-in-rishikesh/26-km-marine-drive-to-nim-beach", "26 km rafting rishikesh"),
    ("/river-rafting-in-rishikesh/36-km-kaudiyala-to-nim-beach", "36 km rafting rishikesh"),
])

subheading("4.4  Adventure Packages & Combos  (12)")
url_table([
    ("/rishikesh-adventure-packages", "rishikesh adventure packages"),
    ("/rafting-and-camping-package-rishikesh", "rafting and camping rishikesh"),
    ("/zipline-rafting-camping-package-rishikesh", "zipline rafting camping combo"),
    ("/bungee-rafting-camping-combo-rishikesh", "bungee rafting camping combo"),
    ("/camping-and-kayaking-course-rishikesh", "kayaking course rishikesh"),
    ("/rishikesh-1-night-2-day-package", "rishikesh 1 night 2 day package"),
    ("/rishikesh-2-night-3-day-package", "rishikesh 2 night 3 day package"),
    ("/weekend-adventure-package-rishikesh", "rishikesh weekend trip package"),
    ("/family-adventure-package-rishikesh", "family adventure trip rishikesh"),
    ("/corporate-team-outing-rishikesh", "corporate team outing rishikesh"),
    ("/college-group-trip-rishikesh", "college group trip rishikesh"),
    ("/delhi-to-rishikesh-adventure-tour", "delhi to rishikesh adventure tour"),
])
doc.add_page_break()

subheading("4.5  Bike & Scooty Rentals — Hub & Model Pages  (9, NEW)")
url_table([
    ("/bike-rental-in-rishikesh", "bike rental in rishikesh"),
    ("/scooty-rental-in-rishikesh", "scooty rental in rishikesh"),
    ("/bike-rental-in-rishikesh/royal-enfield-classic-350", "royal enfield rental rishikesh"),
    ("/bike-rental-in-rishikesh/royal-enfield-himalayan", "himalayan rental rishikesh"),
    ("/bike-rental-in-rishikesh/ktm-duke-390", "ktm rental rishikesh"),
    ("/scooty-rental-in-rishikesh/honda-activa", "activa on rent rishikesh"),
    ("/scooty-rental-in-rishikesh/tvs-jupiter", "scooter on rent rishikesh"),
    ("/atv-quad-biking-in-rishikesh", "atv quad biking rishikesh"),
    ("/bike-rental-price-in-rishikesh", "bike rent price rishikesh"),
])

subheading("4.6  Bike-Tour Routes — Himalayan Circuit  (6)")
url_table([
    ("/bike-tour-rishikesh-to-kheerganga", "rishikesh to kheerganga bike trip"),
    ("/bike-tour-rishikesh-to-chopta-tungnath", "rishikesh to chopta bike trip"),
    ("/bike-tour-rishikesh-to-kedarnath", "rishikesh to kedarnath bike trip"),
    ("/bike-tour-rishikesh-to-badrinath", "rishikesh to badrinath bike trip"),
    ("/bike-tour-rishikesh-to-auli", "rishikesh to auli bike trip"),
    ("/bike-tour-rishikesh-to-spiti-valley", "rishikesh to spiti bike trip"),
])

subheading("4.7  Char Dham & Yatra from Rishikesh  (7, NEW)")
url_table([
    ("/char-dham-yatra-from-rishikesh", "char dham yatra from rishikesh"),
    ("/do-dham-yatra-from-rishikesh", "do dham yatra from rishikesh"),
    ("/kedarnath-tour-from-rishikesh", "kedarnath tour from rishikesh"),
    ("/badrinath-tour-from-rishikesh", "badrinath tour from rishikesh"),
    ("/gangotri-tour-from-rishikesh", "gangotri tour from rishikesh"),
    ("/yamunotri-tour-from-rishikesh", "yamunotri tour from rishikesh"),
    ("/char-dham-by-bike-from-rishikesh", "char dham by bike rishikesh"),
])

subheading("4.8  Haridwar Tours  (4, NEW)")
url_table([
    ("/haridwar-tour-from-rishikesh", "haridwar tour from rishikesh"),
    ("/haridwar-ganga-aarti-tour", "haridwar ganga aarti tour"),
    ("/things-to-do-in-haridwar", "things to do in haridwar"),
    ("/rishikesh-haridwar-combo-tour", "rishikesh haridwar combo tour"),
])

subheading("4.9  Nearby Tour Add-ons  (4)")
url_table([
    ("/mussoorie-tour-from-rishikesh", "mussoorie tour from rishikesh"),
    ("/dhanaulti-tour-from-rishikesh", "dhanaulti tour from rishikesh"),
    ("/nainital-tour-from-rishikesh", "nainital tour from rishikesh"),
    ("/jim-corbett-tour-from-rishikesh", "jim corbett tour from rishikesh"),
])

subheading("4.10  Supporting SEO Guides  (7)")
url_table([
    ("/rishikesh-adventure-sports-price-list", "rishikesh adventure sports price"),
    ("/best-time-for-river-rafting-in-rishikesh", "best time for rafting rishikesh"),
    ("/how-to-reach-rishikesh", "how to reach rishikesh"),
    ("/where-to-stay-in-rishikesh", "where to stay in rishikesh"),
    ("/rishikesh-travel-guide", "rishikesh travel guide"),
    ("/adventure-sports-safety-in-rishikesh", "rishikesh adventure safety"),
    ("/things-to-do-in-rishikesh", "things to do in rishikesh"),
])

subheading("4.11  Legal / Utility  (4)")
url_table([
    ("/privacy-policy", "—"),
    ("/terms-and-conditions", "—"),
    ("/cancellation-and-refund-policy", "—"),
    ("/sitemap.xml  &  /robots.txt", "crawlability / indexing"),
])

# total strip
tt = doc.add_table(rows=1, cols=1); tt.alignment = WD_TABLE_ALIGNMENT.CENTER
cc = tt.rows[0].cells[0]; shade_cell(cc, GOLD_HEX); set_cell_margins(cc, 90, 90, 140, 140); cc.width = Inches(6.3)
p = cc.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; no_space(p)
r = p.add_run("TOTAL: 89 SEO-optimised pages at launch"); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = DARK

# ============================================================ 5. SEO + BACKLINKS
heading("5.  SEO & Backlink Strategy")
subheading("URL & on-page principles")
bullet(" lowercase, hyphen-separated, keyword-first, no IDs or query strings, max 2 levels deep.", "Clean slugs — ")
bullet(" one primary keyword per page; hub pages target broad terms, sub-pages target long-tail.", "Intent mapping — ")
bullet(" Service, Product, FAQPage, BreadcrumbList, TouristTrip & Review schema on every page (reusing the parent site's schema system).", "Structured data — ")
bullet(" hub pages link to sub-pages and combos; combos link back; yatra & bike tours cross-link to cab transfers on the main domain.", "Internal linking — ")
subheading("Backlinks & off-page (Phase 1 scope)")
bullet(" reciprocal links between trivenicabs.in and adventure.trivenicabs.in (cab ↔ adventure / yatra) to pass authority to the new subdomain.", "Cross-domain backlinks — ")
bullet(" create and verify a dedicated Google Business Profile for the adventure vertical (NAP, photos, services, reviews).", "Google Business Profile — ")
bullet(" register adventure.trivenicabs.in as a new Search Console property and submit its sitemap.xml.", "Search Console — ")
body("Further off-page activity is planned for a later phase once the site is live and indexed.", after=4, color=GREY, size=9.5)

# ============================================================ 6. CONTENT
heading("6.  Per-Page Content Blueprint")
body("Each activity / package / tour / rental page follows a proven, conversion-first template:", after=4)
bullet(" Hero with price-from, rating, instant Call / WhatsApp / Book CTA.")
bullet(" Direct-answer intro paragraph (voice-search & featured-snippet ready).")
bullet(" Details table: price, duration, distance/height, age & weight limits, location, inclusions.")
bullet(" What's included / excluded, what to carry, safety & guidelines.")
bullet(" Photo gallery, customer reviews, related activities & combos.")
bullet(" FAQ section (8–14 Q) with FAQ schema; sticky mobile booking bar.")

# ============================================================ 7. TECH
heading("7.  Technical Plan")
bullet(" Next.js 15 + React 19 + Tailwind (App Router) — same stack as the main site for shared components & speed.", "Stack — ")
bullet(" adventure.trivenicabs.in via DNS CNAME / A-record; SSL auto-provisioned.", "Subdomain — ")
bullet(" Separate Next.js deployment (or route group) with ISR (revalidate hourly) for fast, fresh pages.", "Hosting — ")
bullet(" WhatsApp + tel: deep links and an enquiry form; payment gateway optional in a later phase.", "Booking — ")
bullet(" GA4 + WhatsApp click tracking + Search Console, mirroring the parent site's analytics.", "Analytics — ")

# ============================================================ 8. TIMELINE
heading("8.  Phased Rollout & Timeline")
t = doc.add_table(rows=1, cols=4); t.style = "Table Grid"
hdr = ["Phase", "Scope", "Key Pages", "Est. Time"]; wd = [Inches(1.3), Inches(2.7), Inches(1.5), Inches(0.9)]
for i, h in enumerate(hdr):
    cc = t.rows[0].cells[i]; shade_cell(cc, DARK_HEX); set_cell_margins(cc); cc.width = wd[i]
    p = cc.paragraphs[0]; no_space(p); r = p.add_run(h); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = WHITE
phases = [
    ("Phase 1\nFoundation", "Subdomain, design system, homepage, core/conversion pages, analytics, schema, GBP & cross-domain backlinks", "Core + Home (6)", "Week 1–2"),
    ("Phase 2\nAdventure Core", "All activity hubs + rafting stretch pages + top combo packages", "~42 pages", "Week 3–5"),
    ("Phase 3\nYatra & Rentals", "Char Dham/Haridwar/nearby tours + bike & scooty rentals + bike-tour routes", "~30 pages", "Week 5–7"),
    ("Phase 4\nContent & SEO", "Guides, price list, internal linking, sitemap & GSC submission", "~11 pages", "Week 7–8"),
    ("Phase 5\nLaunch & Grow", "Cross-domain links live, reviews, ads-ready landing pages, ongoing optimisation", "Ongoing", "Week 8+"),
]
for idx, row in enumerate(phases):
    cells = t.add_row().cells
    if idx % 2 == 1:
        for ci in range(4): shade_cell(cells[ci], "F4F6F8")
    for ci, val in enumerate(row):
        set_cell_margins(cells[ci]); cells[ci].width = wd[ci]
        p = cells[ci].paragraphs[0]; no_space(p); r = p.add_run(val); r.font.size = Pt(9)
        if ci == 0: r.bold = True; r.font.color.rgb = TEAL
        else: r.font.color.rgb = SLATE

# ============================================================ 9. DELIVERABLES
heading("9.  Deliverables & Next Steps")
subheading("What you receive")
bullet(" A live, mobile-first adventure.trivenicabs.in with 89 SEO-optimised pages.")
bullet(" Full schema, sitemap, robots, analytics, Search Console and Google Business Profile setup.")
bullet(" Cross-domain backlinks connecting the cab and adventure brands.")
bullet(" Booking flow via WhatsApp/call/form (online payments optional later).")
bullet(" Content for all activities, packages, yatra/Haridwar tours, rentals and supporting guides.")
subheading("To proceed, we need from you")
bullet(" Final activity, tour & rental pricing and inclusions.")
bullet(" Vendor/operator tie-ups for rafting, bungee, camping and the rental fleet.")
bullet(" Photos/videos and the booking phone/WhatsApp number for the new vertical.")
bullet(" Sign-off on this plan to begin Phase 1.")

sp = doc.add_paragraph(); sp.paragraph_format.space_before = Pt(10)
ct = doc.add_table(rows=1, cols=1); ct.alignment = WD_TABLE_ALIGNMENT.CENTER
cc = ct.rows[0].cells[0]; shade_cell(cc, GOLD_HEX); set_cell_margins(cc, 140, 140, 160, 160); cc.width = Inches(6.3)
p = cc.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; no_space(p)
r = p.add_run("Ready to launch India's adventure capital under the Triveni brand.")
r.bold = True; r.font.size = Pt(12); r.font.color.rgb = DARK
p2 = cc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER; no_space(p2)
r = p2.add_run("Approve this proposal to begin Phase 1 — adventure.trivenicabs.in")
r.font.size = Pt(10); r.font.color.rgb = DARK

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(14)
r = p.add_run("Triveni Cabs  •  Adventure Vertical Proposal v1.0  •  Confidential  •  June 2026")
r.font.size = Pt(8); r.italic = True; r.font.color.rgb = GREY

out = r"C:\Users\bolis\Desktop\Triveni_Adventure_Rishikesh_Launch_Plan.docx"
doc.save(out)
print("SAVED:", out)
