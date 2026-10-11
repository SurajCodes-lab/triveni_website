"""Reproducible local evidence audit and organic Week 51 planning pack.
No production website files are modified. Editorial review remains explicit.
"""
from pathlib import Path
from urllib.parse import urlsplit
from html.parser import HTMLParser
from collections import Counter
import json, re, csv, hashlib, shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent.parent
OUT=Path(r'D:/Desktop/triveni_next/2.0/week51'); OUT.mkdir(parents=True,exist_ok=True)
D=json.loads((OUT/'source-data.json').read_text(encoding='utf-8'))
APP=ROOT/'.next/server/app'
SITE='https://www.trivenicabs.in'
def pathof(u): return urlsplit(u).path.rstrip('/') or '/'
sitemap={pathof(u) for u in re.findall(r'<loc>(.*?)</loc>',(APP/'sitemap.xml.body').read_text(encoding='utf-8'))}
known=set(json.loads((ROOT/'src/generated/known-routes.json').read_text(encoding='utf-8')))
redirects={r['source']:r['destination'] for r in json.loads((ROOT/'.next/routes-manifest.json').read_text())['redirects'] if ':' not in r['source'] and not r.get('has')}
def resolve(p):
    seen=set()
    while p in redirects and p not in seen:
        seen.add(p); p=redirects[p]
    return p
gsc={pathof(r['Top pages']):r for r in D['Pages']}
htmlfiles={('/'+str(p.relative_to(APP)).replace('\\','/')[:-5]).replace('/index','/',1) if p==APP/'index.html' else '/'+str(p.relative_to(APP)).replace('\\','/')[:-5]:p for p in APP.rglob('*.html')}
excluded={p for p in known|set(htmlfiles)|sitemap|set(gsc) if p.startswith('/_') or re.search(r'\.(png|svg|ico|jpg|jpeg|webp|xml|txt|json|woff2?|css|js)$',p)}
paths=sorted((known|set(htmlfiles)|sitemap|set(gsc))-excluded)

class Meta(HTMLParser):
    def __init__(self):
        super().__init__();self.title=[];self.h1=[];self.canonical='';self.description='';self.robots='';self.links=[];self.tag='';self.ignore=0
    def handle_starttag(self,t,a):
        a=dict(a)
        if t in ('script','style'):self.ignore+=1
        if t in ('title','h1'):self.tag=t
        if t=='meta' and a.get('name','').lower() in ('description','robots'):setattr(self,a['name'].lower(),a.get('content',''))
        if t=='link' and a.get('rel')=='canonical':self.canonical=a.get('href','')
        if t=='a' and a.get('href','').startswith('/'):self.links.append(a['href'])
    def handle_endtag(self,t):
        if t in ('script','style'):self.ignore=max(0,self.ignore-1)
        if t==self.tag:self.tag=''
    def handle_data(self,s):
        if self.tag and not self.ignore:getattr(self,self.tag).append(s.strip())

def family(p):
    if p.startswith('/distance/'):return 'Distance'
    if p.startswith('/toll-cost/'):return 'Toll'
    if p.startswith('/blog/'):return 'Guide'
    if any(x in p for x in ('airport','station-taxi')):return 'Airport / station'
    if any(x in p for x in ('places-to-visit','weekend-getaway')):return 'Weekend / nearby'
    if any(x in p for x in ('tempo-traveller','urbania','maharaja-tempo')):return 'Group vehicle'
    if p.startswith(('/one-way-cab/','/round-trip/')) or (p.count('/')==1 and '-to-' in p):return 'Booking / corridor'
    return 'Other / manual intent review'

actions={
 'Distance':'Check this route\'s endpoints, realistic driving range, stop choices and onward booking destination; preserve Week 50 additions and avoid repeating its toll page.',
 'Toll':'Verify the actual route and applicable toll classes against current official evidence; explain inclusions and uncertainty, then link the relevant fare and booking pages.',
 'Guide':'Verify attraction-specific opening rules, entry conditions, time required and pickup access; answer the leading query early and add an appropriate onward journey link.',
 'Weekend / nearby':'Validate destination choices by travel time, traveller type and overnight suitability; strengthen original decision advice and relevant trip links.',
 'Group vehicle':'Confirm available vehicle, usable seats, luggage, minimum distance and inclusions with operations; match city intent and provide a clear enquiry route.',
 'Booking / corridor':'Confirm one-way versus return intent, applicable fare and inclusions, pickup details and actual availability; add route-specific booking guidance after query-to-page validation.',
 'Airport / station':'Confirm terminal or station, pickup meeting point, waiting rules and luggage requirements; test the booking destination and preserve selected form context.',
 'Other / manual intent review':'Determine the real visitor task, query set and distinct role of this URL before drafting a page-specific improvement.'}

rows=[]
for p in paths:
    m=Meta();file=htmlfiles.get(p)
    if file:m.feed(file.read_text(encoding='utf-8'))
    s=gsc.get(p);f=family(p);dest=resolve(p)
    issues=[]
    if dest!=p:issues.append('Configured redirect: review destination; do not rewrite alias')
    if not file:issues.append('No local built HTML: live verification needed')
    else:
        if not m.title:issues.append('No title found in local HTML')
        if not m.h1:issues.append('No H1 text found in local HTML')
        if not m.description:issues.append('No description found in local HTML')
        if not m.canonical:issues.append('No canonical found in local HTML')
        elif pathof(m.canonical)!=p:issues.append('Canonical points to another path: inspect intent')
        if 'noindex' in m.robots.lower():issues.append('Noindex directive: verify intended purpose')
    bad=sorted({pathof(x) for x in m.links if pathof(x) not in known and pathof(x) not in htmlfiles and pathof(x) not in redirects and not re.search(r'\.[a-z0-9]+$',pathof(x)) and not pathof(x).startswith('/api/')})
    if bad:issues.append(f'{len(bad)} internal paths absent from local inventory; verify live')
    c=s['Clicks'] if s else None; imp=s['Impressions'] if s else None
    priority='P3 - evidence needed'
    if s:priority='P2 - measured demand'
    if s and (imp>=40000 or c>=400):priority='P1 - traffic / protect'
    if s and f in ('Booking / corridor','Group vehicle','Airport / station') and imp>=500:priority='P1 - commercial review'
    if dest!=p:priority='Alias - destination review'
    rows.append({'Page ID':'W51-'+hashlib.sha256(p.encode()).hexdigest()[:12], 'Path':p,'URL':SITE+p,'Family':f,'In local sitemap':p in sitemap,'Built HTML':bool(file),'Configured destination':dest,'Local canonical':m.canonical,'Title':' '.join(m.title),'H1':' | '.join(m.h1),'Description':m.description,'Robots':m.robots,'GSC evidence':'Observed in exported top 1000' if s else 'Unavailable in this truncated export','Clicks':c,'Impressions':imp,'CTR':c/imp if imp else None,'Position':s['Position'] if s else None,'Priority':priority,'Local checks':'; '.join(issues) or 'No issue found by limited metadata checks','Link candidates to verify':' | '.join(bad[:30]),'Next action':actions[f],'Query-page mapping':'Pending filtered GSC export','Editorial review':'Pending individual review','Implementation':'Not started - planning record','Owner role':'SEO + content; operations verifies facts','Verification':'Pending live URL, mobile booking and tracking checks','Outcome':'Pending post-publication observation'})
titles=Counter(r['Title'] for r in rows if r['Title'])
for r in rows:
    r['Same-title URLs']=titles[r['Title']] if r['Title'] else 0
    if r['Same-title URLs']>1:r['Local checks']+='; duplicate title in inventory: inspect aliases before editing'
rows.sort(key=lambda r:(r['Priority'], -(r['Impressions'] or 0),r['Path']))
by_path={r['Path']:r for r in rows}
def stats(items):
    c=sum(x['Clicks'] for x in items);i=sum(x['Impressions'] for x in items)
    return {'Clicks':c,'Impressions':i,'CTR':c/i,'Position':sum(x['Position']*x['Impressions'] for x in items)/i,'Daily clicks':c/len(items)}
periods=[{'Period':name,**stats(items)} for name,items in [('Full 92 days',D['Chart']),('Latest 7 days',D['Chart'][-7:]),('Previous 7 days',D['Chart'][-14:-7]),('Latest 28 days',D['Chart'][-28:]),('Previous 28 days',D['Chart'][-56:-28])]]
base=periods[1];dailyimp=base['Impressions']/7

trafficpaths=['/blog/akshardham-temple-delhi-guide-2026','/blog/ayodhya-ram-mandir-complete-travel-guide-2026','/blog/india-gate-delhi-complete-guide-2026','/blog/hawa-mahal-jaipur-complete-guide-2026','/blog/lotus-temple-delhi-complete-guide-2026','/blog/nahargarh-fort-jaipur-complete-guide-2026','/toll-cost/delhi-to-dehradun','/toll-cost/delhi-to-haridwar']
commercialpaths=['/delhi-to-agra','/chandigarh-to-kasauli','/chandigarh-to-amritsar','/delhi-to-dehradun','/tempo-traveller/chandigarh','/tempo-traveller/lucknow']
briefs=[]
for group,selected in [('Traffic',trafficpaths),('Commercial',commercialpaths)]:
    for p in selected:
        dest=resolve(p);r=by_path.get(dest) or by_path.get(p)
        briefs.append({'Workstream':group,'Requested route':p,'Working destination':dest,'Clicks':gsc.get(dest,{}).get('Clicks'),'Impressions':gsc.get(dest,{}).get('Impressions'),'Position':gsc.get(dest,{}).get('Position'),'Action':actions[family(dest)],'Gate':'Verify live canonical and query-to-page ownership before editing','Status':'Proposed Week 51 batch'})

def write_csv(file,data):
    with file.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(data[0]));w.writeheader();w.writerows(data)
write_csv(OUT/'all-pages-register.csv',rows)
(OUT/'audit-summary.json').write_text(json.dumps({'scope':'Local build snapshot plus GSC export; no live audit performed','inventory_urls':len(rows),'known_routes':len(known),'local_html_files':len(htmlfiles),'sitemap_urls':len(sitemap),'excluded_assets_and_system_paths':len(excluded),'gsc_observed_inventory_urls':sum(r['Clicks'] is not None for r in rows),'families':dict(Counter(r['Family'] for r in rows)),'periods':periods},indent=2),encoding='utf-8')
page_dir=OUT/'page-records';page_dir.mkdir(exist_ok=True)
for r in rows:
    content=[f"# {r['Page ID']} — {r['Path']}",'',f"URL: {r['URL']}",'','## Recorded baseline','']
    for k in ['Family','In local sitemap','Built HTML','Configured destination','Local canonical','Title','H1','GSC evidence','Clicks','Impressions','CTR','Position','Priority','Local checks','Link candidates to verify']:
        value=r[k];value='Unavailable' if value is None else str(value)
        content.append(f'- **{k}:** {value}')
    content+=['','GSC baseline is the full 92-day export, not a current keyword ranking. Local metadata checks are automated; they do not establish live status, indexation, factual accuracy or editorial quality.','','## Individual work to complete','',r['Next action'],'','1. Export this exact page’s queries for the latest 28 days and the preceding 28 days. Record the intended primary query and relevant secondary questions.','2. Read this page in full on mobile and desktop. Write the actual missing answers and supporting source links here; do not assume the family-level action is a confirmed defect.','3. Review sibling URLs for overlapping intent, then choose an additive edit that preserves useful content, booking functions and tracking.','4. Record exact sections changed, factual verification by operations, and relevant internal links.','5. Validate title, canonical, rendered content, mobile navigation, quote selection, and GA4 events without duplicate firing.','6. Log publication, recrawl evidence, and 7/14/28-day matched-period results. A review may conclude that no edit is needed.','','## Status','', '- Editorial review: **Pending**','- Implementation: **Not started**','- Owner: **Unassigned; SEO/content role recommended**','- Evidence and implementation MD: **To be added after this page is actually changed**','- Search result: **Not yet measurable**','']
    record_file=page_dir/(r['Page ID']+'.md')
    if not record_file.exists():record_file.write_text('\n'.join(content),encoding='utf-8')

wb=Workbook();wb.remove(wb.active)
def sheet(name,data):
    ws=wb.create_sheet(name);headers=list(data[0]);ws.append(headers)
    for r in data:ws.append([r.get(k) for k in headers])
    ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='17324D')
    for column in ws.columns:
        key=column[0].value;ws.column_dimensions[column[0].column_letter].width=55 if key in ('URL','Path','Top pages','Top queries','Next action','Action','Local checks') else 23
        for c in column[1:]:
            c.alignment=Alignment(vertical='top',wrap_text=True)
            if key=='CTR':c.number_format='0.00%'
            elif key in ('Position','Daily clicks'):c.number_format='0.00'
    return ws
sheet('Read me',[{'Item':'Purpose','Detail':'Organic growth, keyword ranking improvement and individual page work. Planning and local automated checks only.'},{'Item':'Window','Detail':'2026-06-24 to 2026-09-23; query and page rows cover the whole window.'},{'Item':'Missing data','Detail':'Blank means not in truncated export, not zero traffic. No query-page join is available.'},{'Item':'Status','Detail':'All individual editorial reviews and implementation are pending. Local build is not a live crawl.'},{'Item':'Inventory','Detail':f'{len(rows)} candidate URL records; includes redirects and historical GSC URLs for reconciliation. Not a claim of that many indexable pages.'}])
sheet('Period comparison',periods)
daily=[]
for i,r in enumerate(D['Chart']):daily.append({**r,'7-day average clicks':sum(x['Clicks'] for x in D['Chart'][i-6:i+1])/7 if i>=6 else None})
ws=sheet('Daily performance',daily);chart=LineChart();chart.title='Organic clicks: daily and seven-day average';chart.add_data(Reference(ws,min_col=2,max_col=2,min_row=1,max_row=ws.max_row),titles_from_data=True);chart.add_data(Reference(ws,min_col=6,max_col=6,min_row=1,max_row=ws.max_row),titles_from_data=True);chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=ws.max_row));chart.width=26;chart.height=12;ws.add_chart(chart,'H2')
sheet('All page records',rows)
sheet('GSC queries',D['Queries']);sheet('GSC pages',D['Pages'])
sheet('Week 51 first batch',briefs)
for name in ['Devices','Countries','Search appearance']:sheet(name,D[name])
wb.save(OUT/'week51_analysis_and_page_register.xlsx')

def table(headers,values):return '\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']+['| '+' | '.join(str(c).replace('|',' / ') for c in r)+' |' for r in values])
md=['# Triveni Cabs — Week 51 internal organic growth plan','',
'## 1. Organic growth priorities','',
f'Week 51 focuses on sustainable organic growth, stronger relevant keyword rankings and better individual pages. The latest seven complete export days average **{base["Daily clicks"]:,.0f} clicks/day**. Use this observed baseline to evaluate progress alongside qualified enquiries, useful search visibility and the quality of each completed page.','',
'Use genuine traveller demand, accurate useful information, better search-intent alignment, stronger relevant internal navigation and reliable booking journeys. Do not buy clicks, manufacture reviews, use manipulative links, stuff keywords or publish city-name variations without distinct value. Keep coverage as supporting evidence; organic clicks and qualified enquiries are the outcomes.','',
'Week 51 starts a complete page-by-page programme. It does not label thousands of pages as improved merely because a shared template changed. Every URL needs its own recorded decision, and some pages may correctly need no change.','',
'## 2. Evidence and limits','',
'Source: trivenicabs.in-Performance-on-Search-2026-09-25.xlsx. The filename is September 25, but the actual daily data ends September 23 and begins June 24: 92 days, Web search. Queries and Pages each contain 1,000 exported rows. No query-to-page pairing, recent-period keyword comparison, search demand forecast, index-coverage export or conversion dataset is supplied.','',
f'The property total is **39,162 clicks** and **6,415,279 impressions**. Exported query rows contain **7,996 clicks**, only **{7996/39162:.1%}** of property clicks. The page table contains **38,108 clicks**; page and property aggregation also differ. Never add these dimensions together. Missing query/page rows are unknown, not zero. Search appearance categories may overlap.','',
'A keyword position in this pack is its average across the full export window; it is not its rank today. Page positions do not prove that a particular keyword ranks at that position. Query-to-page ownership must be obtained using GSC filters. The latest Week 50 deployment happened after this data window, so this export cannot establish the effect of that release.','',
'## 3. Performance diagnosis','',
table(['Period','Clicks','Impressions','CTR','Average position','Clicks/day'],[[r['Period'],f"{r['Clicks']:,.0f}",f"{r['Impressions']:,.0f}",f"{r['CTR']:.3%}",f"{r['Position']:.2f}",f"{r['Daily clicks']:.1f}"] for r in periods]),'',
'The latest seven days gained 16.1% in clicks and 20.6% in impressions versus the preceding seven. CTR fell slightly, from 0.696% to 0.670%. Recent growth therefore came with greater visibility; this is not evidence of a CTR improvement. The impression-weighted daily position improved from about 8.95 to 7.95, but changing query mix also affects this figure.','',
'Across 28 days, clicks fell 2.1% and impressions fell 18.9%, although CTR improved. The recovery needs to persist before it can be called sustained growth. Daily totals, seven-day averages and matched 28-day comparisons should remain separate.','',
'Mobile accounts for 32,530 clicks (83.1%); desktop contributes 6,310, and tablet 322. India contributes 36,493 clicks (93.2%). Prioritise mobile readability, pickup clarity and booking usability for Indian journeys. Device differences alone do not prove a technical cause; investigate page and intent mix before assigning one.','',
'## 4. How organic improvement will be evaluated','',
f'The latest seven days average **{dailyimp:,.0f} impressions/day** at **{base["CTR"]:.2%} CTR**. Improve relevant visibility and the usefulness of the search result together. Assess CTR within comparable query, page, position and device groups rather than judging every page against one site-wide number.','',
'Distance, brand, attraction and booking searches behave differently. A route distance answer may be satisfied directly in Google; a stronger page cannot force a click where the visitor has no further need. Prioritise useful journey questions and relevant commercial intent rather than impressions alone.','',
'Evaluate progress using the trailing seven-day average, matched 28-day comparisons, relevant-query visibility and qualified enquiries. If progress stalls, identify the pages and queries responsible before expanding publication volume. Record delivered improvements separately from results that are still awaiting observation.','',
'## 5. Keyword ranking programme','',
'The following are real exported query baselines. Intent groups are rule-based planning labels and need individual review. Position targets are desired movement after implementation and remeasurement, not forecasts. The workbook preserves all 1,000 queries.','',
table(['Query','Clicks','Impressions','Position','Initial objective'],[[q['Top queries'],int(q['Clicks']),int(q['Impressions']),q['Position'],'Protect' if q['Position']<5 else 'Work toward top 5' if q['Position']<=12 else 'Work toward top 10' if q['Position']<=20 else 'Longer-term intent/authority review'] for q in D['Queries'] if q['Top queries'] in {'delhi to agra cab','chandigarh to kasauli cab','chandigarh to amritsar cab','tempo traveller in chandigarh','tempo traveller in lucknow','delhi to dehradun taxi','delhi to jaipur cab','delhi to manali cab','delhi to chandigarh cab','rishikesh taxi service','weekend getaways from noida','nahargarh fort timings','delhi to agra distance'}]),'',
'For each selected query, filter it in GSC, open Pages, and compare the latest 28 days with the preceding 28 days. Record the actual landing URL, country and device. Inspect the search results for dominant intent and features; distinguish official visitor facts, route information and bookable transport. Choose one intended landing page per intent, with supporting pages serving different traveller questions. Similar strings alone do not establish cannibalisation.','',
'Commercial terms around positions 8–12 are the first booking candidates. Terms around 15–20 need a more substantial usefulness and intent review. Terms beyond 20 are a longer-term queue, not quick-win promises. Broad queries such as cab fare or taxi service near me require context; do not prioritise them solely because their text sounds commercial.','',
'## 6. First Week 51 batch: traffic and commercial value','',
'Plan for eight traffic-page reviews and six commercial-page reviews, followed by individually justified edits. These are a proposed quality-first capacity limit, to be adjusted to available reviewers; completion is not assumed. Protect /places-to-visit-near-noida, which already delivers 2,323 clicks in the export. Review any modification to an existing winner against its successful query set.','',
table(['Workstream','Canonical working path','Clicks / impressions','Position'],[[b['Workstream'],b['Working destination'],f"{int(b['Clicks']) if b['Clicks'] is not None else 'Unavailable'} / {int(b['Impressions']) if b['Impressions'] is not None else 'Unavailable'}",b['Position'] if b['Position'] is not None else 'Unavailable'] for b in briefs]),'']
for b in briefs:
    r=by_path.get(b['Working destination'])
    md += [f"### {b['Working destination']}",'',b['Action'],'',f"Local title: {r['Title'] if r else 'Unavailable'}. Local checks: {r['Local checks'] if r else 'Needs live verification'}.",'',
    'First obtain this exact page’s leading queries and compare the live answer with the visitor’s task. Record what is already answered before drafting. Write only the missing, verified detail, make the title accurately describe the page, and add a contextual link to the next useful step. Keep the booking offer appropriate to the informational or commercial intent.','',
    'Acceptance: a reviewer can identify the specific new traveller benefit, verify each changed factual claim, follow the intended canonical destination and complete the relevant mobile enquiry action. Record the before/after sections and subsequent GSC results separately.','']
md += ['## 7. Individual coverage of the entire site','',
f'The reconciled planning register contains **{len(rows):,} candidate URL records** from **{len(known):,} known routes**, **{len(htmlfiles):,} local HTML files**, **{len(sitemap):,} sitemap URLs** and the 1,000-row GSC page export. Assets and internal system paths are excluded from the page register. Redirect aliases and historical URLs remain visible for reconciliation; this is not a claim that every record is a distinct live indexable page.','',
table(['Family','Candidate URLs'],sorted(Counter(r['Family'] for r in rows).items())),'',
'Each record has a stable page ID, its own Markdown work file, exact URL, source presence, observed GSC metrics where available, local title/H1/canonical, limited local-check findings, priority, next action, and separate editorial/implementation/measurement statuses. All editorial reviews are currently pending. Automated checks are completed only to the limited extent described; generating a record is not completing that page’s work.','',
'Start with the 14 selected reviews, then work the P1 register in descending relevant opportunity. Run P2 measured-demand reviews next, and collect evidence for P3 pages omitted from the export. Aliases are resolved to their destinations before editing. A redirect or intentional noindex page needs its own documented decision, not fabricated ranking goals.','',
'Individual completion requires: actual intent and query ownership; a full read; verified unique information; a recorded edit or justified keep decision; internal-link and mobile journey checks; preservation of analytics; and a dated observation record. Shared-template fixes can be applied efficiently but every affected URL still requires a verified outcome. Do not use a shared rewrite as proof of 3,800 individual improvements.','',
'At 20 full editorial reviews per working day, 3,800 pages require 190 reviewer-days; at 40/day they require 95. These are workload arithmetic, not staffing commitments. Deep factual reviews may take longer. Week 51 establishes the register, completes the first batch and measures actual throughput so the remaining sequence can be scheduled honestly. No page is removed because it falls outside the first batch.','',
'## 8. Organic content and authority work','',
'For route pages, improve the decision a traveller actually faces: pickup location, duration variability, a practical stop, luggage fit, route-specific costs and what the quote includes. Ask operations for original, accurate answers and attributable photographs or experience where available. For visitor guides, distinguish verified opening or entry rules from transport advice. Source changing facts before publication and make uncertain estimates clear.','',
'Avoid mechanical word-count expansion and repeated paragraphs with substituted city names. New pages require a demonstrated unanswered intent and a confirmed absence of a suitable existing page. There is no automatic ten-page quota. Strengthening a relevant existing page can be the better decision.','',
'Improve contextual links where they help the reader choose the next step: guide to suitable route, route to verified fare details, vehicle choice to relevant booking. Review competing pages by actual query overlap before proposing consolidation. Existing useful content and functionality stay in place; no mass deletions or redirect changes are part of this plan.','',
'Earn references through genuinely useful original route resources and real business relationships. Keep real business details accurate wherever the business already appears. Outreach can be prepared for review, but no messages are sent as part of this planning task. Success is useful referral and organic discovery, not an arbitrary backlink count.','',
'## 9. Technical support and Week 50 carry-forward','',
'Preserve GA4 in layout.js, Search Console verification, existing analytics events, booking forms, quote state and all useful content. Any performance change must demonstrate that tracking still loads and fires as intended. Measure GA4 enquiry actions separately from GSC clicks: they are different metrics with different collection rules.','',
'The Week 50 records report 96 remaining broken-target cases after earlier repairs. Reproduce these against current output before carrying them into an implementation batch. Investigate the local register’s link candidates using live responses; absence from a generated route list alone is not proof of a broken link.','',
'Check the priority URLs for successful live response, indexability, correct canonical, useful rendered content and mobile usability. Coverage is a diagnostic aid: repair an unintended barrier on a valuable page, but do not treat the indexed-page count as the primary growth target. Do not remove intentional noindex rules without checking page purpose.','',
'Performance work should begin with a measured mobile bottleneck on a representative priority page, preserve consent and tracking behaviour, and be verified after a small change. Structured data must reflect visible, supported content; do not promise stars or introduce invented reviews. Preserve the recent ten-page Week 50 pilot and allow a post-release measurement window before attributing success or failure.','',
'## 10. Week 51 execution sequence','',
table(['Stage','Responsible role','Deliverable and gate'],[
['Day 1','SEO / analyst','Reconcile live canonicals; obtain 28-day comparison and query-to-page exports for first batch; freeze baselines.'],
['Day 2','Content + operations','Individually review first four traffic pages; verify facts; draft additive changes and specific acceptance notes.'],
['Day 3','Content + developer','Review remaining four traffic pages; implement approved factual drafts in small batches; check mobile and tracking.'],
['Day 4','SEO + operations','Review first three commercial pages: actual query ownership, fare interpretation, pickup and vehicle choice.'],
['Day 5','Content + developer','Review remaining three commercial pages and implement justified edits; verify canonical links and booking state.'],
['Day 6','Developer + reviewer','Reproduce high-value link issues, complete targeted fixes and regression checks; prepare per-change Markdown records.'],
['Day 7','SEO / owner','Close evidence gaps, publish the completed-work register, log release and monitoring cohorts, schedule the next P1 batch from measured capacity.']]),'',
'Roles describe responsibilities, not assigned staff. If verification capacity is lower, carry unfinished pages forward explicitly. Do not publish unverified facts simply to meet a daily count. Actual implementation records must distinguish planned, implemented, verified, deployed and measured states.','',
'## 11. Measurement and decisions after publication','',
'Capture a 28-day pre-change page/query baseline and the preceding 28 days. Preserve the deployment date in the internal measurement record. Review crawl and technical behaviour first; then assess 7-, 14- and 28-day observations after publication, allowing for data latency. Early movements can reflect demand, ranking mix and search-result features as well as the edit.','',
'Report property clicks, seven-day average daily clicks, 28-day clicks, non-brand relevant-query clicks where measurable, cohort impressions/CTR/position, and qualified enquiries or bookings where tracking supports them. Do not infer bookings from clicks. Use unchanged comparable pages as a contextual comparison, while acknowledging this is not a randomised experiment.','',
'For each target query, report both query-level position and the intended page receiving impressions. A gain on the wrong page may call for clearer intent separation. If impressions increase but clicks do not, review search-result intent and snippet accuracy. If both decline, investigate demand, indexability and ranking by query rather than blindly changing the title again.','',
'Completion report: planned versus delivered; additional work only where something extra actually happened; exact modified URLs; factual sources; tracking checks; deployment status; pending decisions; and observed results with their time windows. Each implemented page or shared change gets its own Markdown evidence record.','',
'## 12. Files, interpretation and immediate next work','',
'week51_analysis_and_page_register.xlsx contains the baseline, daily chart, full query/page exports, candidate URL register and first batch. all-pages-register.csv is the portable register. page-records contains one initial work record per candidate URL. audit-summary.json preserves computed counts. The Python builder regenerates baseline outputs but preserves existing individual Markdown records. Keep working status changes in those records; manually edited workbook and CSV outputs need to be saved separately before rebuilding.','',
'This delivery prepares the plan and local baseline. It does not deploy website changes or mark the whole site as editorially reviewed. The next implementation work starts with verified GSC query ownership and the first page-specific content review, retaining every existing protection requested for GA4, Search Console and site content.','',
'Guidance: [Google: helpful, reliable, people-first content](https://developers.google.com/search/docs/fundamentals/creating-helpful-content) supports focusing on actual visitor usefulness and original value. [Search Console performance report](https://support.google.com/webmasters/answer/7576553?hl=en) explains the performance dimensions used here. The numerical findings in this plan come from the supplied workbook and the local build, not from these guidance pages.','']
intent_rows=[]
for intent in sorted({q.get('intent','unclassified') for q in D['Queries']}):
    qs=[q for q in D['Queries'] if q.get('intent','unclassified')==intent]
    c=sum(q['Clicks'] for q in qs);im=sum(q['Impressions'] for q in qs)
    intent_rows.append([intent,len(qs),int(c),int(im),f'{c/im:.2%}',f"{sum(q['Position']*q['Impressions'] for q in qs)/im:.2f}"])
cohort=[gsc[p] for p in trafficpaths if p in gsc]
cohort_c=sum(p['Clicks'] for p in cohort);cohort_i=sum(p['Impressions'] for p in cohort)
md += ['## Appendix A. Query intent and traffic economics','',
table(['Exploratory intent','Queries','Clicks','Impressions','CTR','Position'],intent_rows),'',
'These classifications use query words and can contain mixed intent. They describe only the exported rows, which represent about one fifth of all clicks. In particular, high distance-query impressions and low CTR do not establish that a title edit can recover all unclicked impressions. Brand clicks are valuable but should be separated from new non-brand demand when assessing growth.','',
f'The eight selected traffic pages together recorded **{cohort_c:,} clicks** from **{cohort_i:,} impressions** over 92 days, a combined CTR of **{cohort_c/cohort_i:.2%}**. At exactly the same impression volume, an illustrative 1% combined CTR would add **{max(0,cohort_i*.01-cohort_c)/92:.1f} clicks/day** over their historical rate; 2% would add **{max(0,cohort_i*.02-cohort_c)/92:.1f}**. These are sensitivity calculations, not expected CTRs. They deliberately use one consistent 92-day page cohort and do not add query totals to page totals.','',
'The first batch establishes a measured process that can expand across further useful pages and relevant query demand. The stronger commercial pages also serve a different purpose: qualified journey enquiries may matter more than their contribution to gross click volume.','',
'## Appendix B. Next measured-demand queue','',
'After the first batch, validate these high-impression pages individually. A queue entry authorises investigation in the plan; it does not establish a defect or prescribe a rewrite. Distance pages recently changed in Week 50 should first be checked against their implementation records to avoid overlapping edits.','',
table(['URL path','Clicks','Impressions','CTR','Position'],[[r['Path'],int(r['Clicks']),int(r['Impressions']),f"{r['CTR']:.2%}",r['Position']] for r in sorted((r for r in rows if r['Clicks'] is not None and r['Path'] not in trafficpaths and r['Path'] not in {resolve(p) for p in commercialpaths}),key=lambda r:-r['Impressions'])[:25]]),'']
text='\n'.join(md);(OUT/'week51_organic_growth_plan.md').write_text(text,encoding='utf-8')

doc=Document();section=doc.sections[0];section.top_margin=section.bottom_margin=Inches(.8);section.left_margin=section.right_margin=Inches(.85)
style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(10.5);style.paragraph_format.space_after=Pt(7)
for name in ['Title','Heading 1','Heading 2']:doc.styles[name].font.color.rgb=RGBColor.from_string('17324D')
lines=text.splitlines();i=0
while i<len(lines):
    s=lines[i]
    if not s:i+=1;continue
    if s.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].startswith('|'):
            cells=[x.strip() for x in lines[i].strip('|').split('|')]
            if not all(re.fullmatch('-+',x) for x in cells):block.append(cells)
            i+=1
        t=doc.add_table(rows=1,cols=len(block[0]));t.style='Light Shading Accent 1'
        for j,c in enumerate(block[0]):t.rows[0].cells[j].text=c
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for row in block[1:]:
            cells=t.add_row().cells
            for j,c in enumerate(row):cells[j].text=c
        continue
    clean=s.replace('**','')
    if s.startswith('# '):doc.add_heading(clean[2:],0)
    elif s.startswith('## '):doc.add_heading(clean[3:],1)
    elif s.startswith('### '):doc.add_heading(clean[4:],2)
    else:doc.add_paragraph(clean)
    i+=1
footer=section.footer.paragraphs[0];footer.text='TRIVENI CABS | WEEK 51 INTERNAL PLAN | '
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.save(OUT/'week51_organic_growth_plan.docx')
print(json.dumps({'candidate_urls':len(rows),'individual_records':len(list(page_dir.glob('*.md'))),'plan_words':len(text.split()),'first_batch':len(briefs),'outputs':str(OUT)},indent=2))
