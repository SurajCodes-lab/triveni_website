"""Build the Week 52 planning pack outside both website repositories."""
from pathlib import Path
import csv, json, re, tempfile
from urllib.parse import urlsplit
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document
from docx.shared import Inches, Pt

ROOT=Path(__file__).resolve().parent.parent
OUT=Path('D:/Desktop/triveni_next/2.0/week52')
OLD=OUT.parent/'week51'
OUT.mkdir(parents=True,exist_ok=True)
evidence_path=OUT/'planning-input-evidence.json'
if not evidence_path.exists():
    evidence_path.write_text((Path(tempfile.gettempdir())/'triveni-week52-evidence.json').read_text(encoding='utf-8'),encoding='utf-8')
E=json.loads(evidence_path.read_text(encoding='utf-8'))
D=json.loads((OLD/'source-data.json').read_text(encoding='utf-8'))
known=set(E['known']); adventure=set(E['adventureRoutes'])
pages={urlsplit(r['Top pages']).path.rstrip('/') or '/':r for r in D['Pages']}
def table(h,rows):return '\n'.join(['| '+' | '.join(h)+' |','| '+' | '.join(['---']*len(h))+' |']+['| '+' | '.join(str(x).replace('|',' / ') for x in r)+' |' for r in rows])
def family(p):
    if p.startswith('/blog/'):return 'Blogs'
    if 'wedding' in p or 'baraat' in p:return 'Wedding'
    if any(x in p for x in ['chardham','char-dham','do-dham','kedarnath','badrinath']):return 'Char Dham / pilgrimage'
    if 'tempo' in p or 'urbania' in p:return 'Tempo / Urbania / Maharaja'
    if p.startswith('/one-way-cab'):return 'One-way cabs'
    if p.startswith('/round-trip'):return 'Round-trip cabs'
    if p.startswith('/distance/'):return 'Distance'
    if p.startswith('/toll-cost/'):return 'Toll'
    if 'airport' in p or 'station' in p:return 'Airport / railway'
    if p.count('/')==1 and '-to-' in p:return 'Intercity routes'
    if p in citypaths:return 'City hubs'
    if any(x in p for x in ['cab-service','local-taxi','outstation','taxi-fare']):return 'Local / outstation / fare'
    return 'Other existing pages'
citypaths=['/'+x['name'].lower().replace(' ','-') for x in E['cities']]
blogs=[
('B01','Lucknow airport or Charbagh pickup: details to share before booking','lucknow-airport-charbagh-cab-pickup-checklist','Lucknow pickup planning',['/airport-service/lucknow','/lucknow'],'Compare pickup information, train or flight details, luggage, meeting-point confirmation and waiting terms. Verify station names and current airport pickup instructions; do not invent terminal rules.'),
('B02','Lucknow to Ayodhya: plan the cab around darshan and the return','lucknow-ayodhya-darshan-cab-day-plan','Lucknow–Ayodhya planning',['/lucknow-to-ayodhya','/distance/lucknow-to-ayodhya'],'Explain one-way drop versus retained vehicle, time buffers, walking needs and return commitments. Keep transport separate from temple passes. Avoid repeating the existing general Ayodhya guide.'),
('B03','Lucknow to Varanasi: when to choose a one-way or return cab','lucknow-varanasi-one-way-or-return-cab','Lucknow–Varanasi booking decision',['/lucknow-to-varanasi','/distance/lucknow-to-varanasi'],'Use traveller scenarios: relocation, a short visit and a multi-day stay. Compare charge components without inventing package prices; cover hotel access and final return pickup.'),
('B04','Lucknow wedding guest transport: airport, hotel and venue pickup plan','lucknow-wedding-guest-transport-pickup-plan','Lucknow wedding logistics',['/wedding/lucknow','/tempo-traveller/lucknow'],'Provide a usable guest-arrival checklist, vehicle allocation and trip sequence. Distinguish the bridal car from guest shuttles. City-specific arrangements must add value beyond the existing wedding transport playbook.'),
('B05','Choosing a tempo traveller in Lucknow: seats, bags and multiple pickups','lucknow-tempo-traveller-seats-luggage-pickups','Lucknow group vehicle choice',['/tempo-traveller/lucknow','/12-seater-vs-17-seater-tempo'],'Explain passenger count versus luggage space, pickup consolidation, actual seat layout and written inclusions. Confirm available configurations with operations; do not use assumed suitcase capacities.'),
('B06','Delhi to Agra cab: direct drop, sightseeing day or overnight return','delhi-agra-cab-drop-sightseeing-overnight','Delhi–Agra trip-type comparison',['/delhi-to-agra','/distance/delhi-to-agra'],'Offer a decision table and quote-request checklist. Clarify waiting, local stops, monument entry and return transport. Differentiate this article from existing itinerary and fare pages.'),
('B07','Chandigarh to Manali group travel: quote and pickup checklist','chandigarh-manali-group-travel-quote-checklist','Chandigarh–Manali group planning',['/chandigarh-to-manali','/tempo-traveller/chandigarh'],'Cover luggage, pickup locations, overnight itinerary, planned breaks and the return leg. Verify operating assumptions with the team; avoid fixed arrival promises or unsupported road-condition statements.'),
('B08','Char Dham road journey: arrival buffers, luggage and vehicle handovers','char-dham-road-trip-arrival-luggage-handover','Char Dham transport preparation',['/chardham-yatra','/char-dham-yatra-by-tempo-traveller'],'Focus on practical coordination rather than another generic cost guide. Verify current opening/closing dates, registration and route advisories before including them; do not imply year-round access.'),
('B09','Delhi to Rishikesh rafting weekend: coordinate the cab and activity slot','delhi-rishikesh-rafting-cab-activity-timing','Rafting arrival and transfer planning',['/delhi-to-rishikesh','/rishikesh'],'Give separate transport and activity checklists; explain confirmation of meeting point, activity timing and return pickup. Link to the Adventure rafting page where that is the next useful action, without claiming a guaranteed slot.'),
('B10','Lucknow to Rishikesh group trip: plan transport, camping and return','lucknow-rishikesh-group-camping-travel-plan','Lucknow adventure group itinerary',['/lucknow','/tempo-traveller/lucknow'],'Plan a realistic multi-day journey with group pickup, luggage and accommodation coordination. Distinguish camp check-in from travel time and confirm the operator’s actual package inclusions. Link to the Adventure camping page.'),
]
blogrows=[]
for ident,title,slug,intent,links,brief in blogs:
    assert '/blog/'+slug not in known
    blogrows.append({'ID':ident,'Proposed title':title,'Proposed URL':'/blog/'+slug,'Intent':intent,'Candidate booking links':' | '.join(p for p in links if p in known),'Brief':brief,'Overlap review':'Required: exact slug absent; semantic overlap not yet cleared','Status':'Planned - not written or published'})

routepaths=['/lucknow-to-ayodhya','/lucknow-to-varanasi','/lucknow-to-prayagraj','/lucknow-to-delhi','/delhi-to-lucknow','/delhi-to-agra','/delhi-to-jaipur','/delhi-to-dehradun','/delhi-to-haridwar','/delhi-to-rishikesh','/chandigarh-to-manali','/chandigarh-to-kasauli','/chandigarh-to-amritsar']
manifest=json.loads((ROOT/'.next/routes-manifest.json').read_text(encoding='utf-8'))
redirects={}
for r in manifest['redirects']:
    if not r.get('has') and not re.search(r'[:*()]',r['source']):redirects.setdefault(r['source'],r['destination'])
def canonical(p):
    seen=set()
    while p in redirects and p not in seen:seen.add(p);p=redirects[p]
    return p
routes=[]
for p in routepaths:
    dest=canonical(p);r=pages.get(dest,{})
    routes.append({'Requested path':p,'Working canonical':dest,'Known locally':dest in known,'Historical clicks':r.get('Clicks'),'Historical impressions':r.get('Impressions'),'Historical position':r.get('Position'),'Next action':'Confirm query ownership; review pickup, trip type, fare inclusions, relevant links and mobile LCP','Status':'Planned; latest-period evidence required'})
cityrows=[]
for p in citypaths:
    r=pages.get(p,{})
    cityrows.append({'Path':p,'Known locally':p in known,'Historical clicks':r.get('Clicks'),'Historical impressions':r.get('Impressions'),'Historical position':r.get('Position'),'Priority':'P1 Lucknow' if p=='/lucknow' else 'P2 all-city review','Content review':'Pending individual review','Mobile LCP':'Not measured','Desktop LCP':'Not measured','Action':'Review actual pickup coverage, routes, service distinctions, verified facts, title and relevant internal links'})

workstreams=[
('Lucknow','City hub plus wedding, tempo, local, one-way, round-trip, airport and core corridor pages','Correct intent, useful pickup details, service-specific FAQs, route selection and contextual navigation; retain successful wedding/tempo answers.'),
('All city hubs',f'All {len(citypaths)} configured city hubs','Give each city an individual review record. Inspect distinct local content, pickup details and route data; do not mass-rewrite city names or claim every page needs editing.'),
('Top routes','13 initial corridors, with canonical resolution','Review actual one-way/return intent, vehicle selection, written inclusions, journey logistics and supporting distance/toll links. Expand from fresh demand evidence.'),
('One-way cabs','All existing one-way URLs; priority corridors first','Make drop intent, pickup, actual price basis and exclusions clear. Preserve booking selection and enquiry context. Check overlaps with root corridor pages.'),
('Round-trip cabs','All existing round-trip URLs; priority corridors first','Clarify billable days/distance, overnight requirements, waiting, stops and return point. Verify the displayed quote changes correctly with trip selection.'),
('Local and outstation cabs','Existing city service, local taxi, outstation and fare families','Differentiate local packages, intercity drops and retained vehicles. Preserve valid package rules; correct contradictions only with verified evidence.'),
('Wedding','Hub, city and wedding vehicle pages; Lucknow first','Separate bridal vehicle and guest transport intent; validate decoration scope, rental duration, waiting and event sequence. Preserve strong pages while improving missing answers.'),
('Char Dham / pilgrimage','Hub, dham, route, package and vehicle pages','Review starting point, transfer versus tour, baggage, itinerary buffers, registration and seasonal status. Current official guidance is required for changing operational claims.'),
('Tempo / Urbania / Maharaja','City hubs, corridors and vehicle comparison pages','Verify seating, bags, real configuration, billing basis and trip suitability; connect equivalent vehicle choices without generating duplicate route pages.'),
('Airport / railway','Existing transfer and station pages','Check exact terminal/station naming, pickup instructions, waiting terms, luggage and return links. Keep selected journey context through enquiry.'),
('Distance / toll / sightseeing / other','Existing inventory retained in the full-site queue','Inspect useful facts, actual route relevance, canonical links and mobile layout. Coverage is supporting evidence, not a substitute for ranking or enquiry results.'),
('Adventure redesign','Homepage, shared category/product layouts, enquiry journey and existing site pages','Improve activity selection, comparison, verified product facts, availability enquiry, brand clarity and fast first paint; retain all existing URLs and useful content.'),
]

cross=[
('/delhi-to-rishikesh','/river-rafting-in-rishikesh','Traveller arranging rafting after arrival','Explore rafting options in Rishikesh'),
('/rishikesh','/camping-in-rishikesh','Visitor choosing an overnight activity stay','Explore camping options'),
('/tempo-traveller/delhi-to-rishikesh','/college-group-trip-rishikesh','Group arranging activity plans','Plan a Rishikesh group outing'),
('/chardham-yatra','/char-dham-yatra-from-rishikesh','Visitor starting a journey from Rishikesh','Explore Char Dham trips from Rishikesh'),
]
crossrows=[]
for parent,child,reason,anchor in cross:
    crossrows.append({'Main-site path':parent,'Adventure path':child,'Main path locally known':parent in known,'Adventure path locally built':child in adventure,'Reader reason':reason,'Example main-to-Adventure anchor':anchor,'Return link purpose':'Relevant cab or transfer options; use the actual final canonical','Status':'Candidate; check live destinations and existing links before adding'})
inventory=[]
with (OLD/'all-pages-register.csv').open(encoding='utf-8-sig',newline='') as f:
    for r in csv.DictReader(f):
        p=r['Path'];inventory.append({'Site':'Triveni Cabs','Path':p,'Family':family(p),'Page ID':r['Page ID'],'Known locally now':p in known,'Priority':'P1 Lucknow' if 'lucknow' in p else 'P1 route' if p in routepaths else 'P2 family review','Review':'Pending individual Week 52 review','Implementation':'Not started','LCP':'Not measured','Previous record':'../week51/page-records/'+r['Page ID']+'.md'})
for p in sorted(adventure):inventory.append({'Site':'Triveni Adventure','Path':p,'Family':'Adventure','Page ID':'ADV-'+p.strip('/').replace('/','--') if p!='/' else 'ADV-home','Known locally now':True,'Priority':'P1 shared redesign' if p in ['/','/book-now','/river-rafting-in-rishikesh','/camping-in-rishikesh'] else 'P2 individual verification','Review':'Pending individual Week 52 review','Implementation':'Not started','LCP':'Not measured','Previous record':''})

md=['# Triveni — Week 52 implementation plan','',
'## 1. Agreed scope','',
'Create exactly **10 new blogs total** this week, published under the main site’s existing /blog/ path. Focus on Lucknow, high-value corridors and all city pages. Include wedding, Char Dham, tempo travellers, Urbania/Maharaja, local and outstation cabs, one-way and round-trip cabs, airport/railway transfers and supporting content. The phrase “two cabs” is interpreted as two-way/round-trip bookings.','',
'Also redesign adventure.trivenicabs.in using D:/adventure_triveni, and add useful contextual links between the two sites. No separate blog subdomain or independent outreach programme is included. The ten-blog count covers the whole weekly publishing programme, not ten per site or per family.','',
'Retain existing pages, useful content, prices unless verified corrections are required, booking functions, GA4 and Search Console verification. Every implementation gets a Markdown record with affected URLs, reason, exact changes, tests and pending items. This document is a plan; it does not claim implementation, deployment, current ranking gains or completed LCP tests.','',
'## 2. Evidence available and Week 51 carry-forward','',
'The inspected Week 51 records include the completed deployment record, restored WhatsApp intent measurement, link repairs, site-wide moderation and measured page records. A later record also describes built content improvements for pilgrimage and city pages. Confirm the actual live release before treating every later built change as deployed. Do not repeat already-completed work or overwrite later edits from an old baseline.','',
'Search data currently available to this plan covers June 24–September 23, 2026. It is historical context, not the latest weekly result. Query and page sheets each contain only 1,000 rows and lack query-to-page comparisons. Obtain the latest 28 days versus the preceding 28, with page/query filters, before final keyword assignments. Missing rows are unavailable rather than zero traffic.','',
f'The current city data defines **{len(citypaths)} city hubs**, all found in the local route inventory. The Adventure build manifest contains **{len(adventure)} non-internal generated routes**; this includes auxiliary routes and is not a live indexation count. The combined work register retains all {len(inventory):,} candidate records from the previous main-site register plus this Adventure snapshot. Validate indexability and route purpose before setting SEO goals for individual records.','',
'The live Adventure URL could not be retrieved by the web tool during this planning inspection. Design observations are from local source, not a visual review of the live site. No Lighthouse or CrUX measurements have been taken in this planning pass.','',
'## 3. Order of work and completion rules','',
table(['Priority','Work','Completion evidence'],[['P0','Freeze current releases and tracking; gather fresh GSC; baseline LCP','Saved evidence, owner, source window and known limitations'],['P1','Lucknow, the initial route list and first city reviews','Individual page decisions and verified additions/corrections'],['P1','Ten article briefs and publication queue','Unique purpose, full editorial review, source checks and working links'],['P1','Adventure homepage/product/enquiry redesign','Mobile and desktop review, correct brand contacts and tested enquiries'],['P2','All remaining city and service-family queues','Each page reviewed or explicitly pending; verified shared fixes per affected URL'],['Release','Build, regression and controlled deployment','Per-change MD, production URL checks, tracking and monitoring records']]),'',
'All families are in scope. This does not mean thousands of full manual rewrites are assumed to fit into one week. Review every city hub individually, run site-wide technical checks, implement high-priority and shared fixes, and retain a visible per-page editorial queue for the rest. A shared template change is not proof that all affected pages received individual factual review.','',
'## 4. Lucknow first','',
table(['Existing page','Historical clicks','Impressions','Average position'],[[p,int(pages[p]['Clicks']),int(pages[p]['Impressions']),pages[p]['Position']] for p in ['/lucknow','/wedding/lucknow','/tempo-traveller/lucknow','/distance/lucknow-to-ayodhya','/distance/lucknow-to-prayagraj','/places-to-visit-near-lucknow'] if p in pages]),'',
'The city hub is a ranking-review priority while wedding and tempo already have meaningful visibility. Protect useful content on the stronger pages. Map city-cab, wedding-car, group-vehicle and corridor intent separately rather than making all pages compete for the same phrase. These positions are full-window historical averages, not ranks today.','',
'Review the Lucknow hub’s actual pickup coverage, airport and station connections, route choices and distinction between local hire, outstation travel and one-way service. Confirm facts with operations. Improve titles only when they accurately describe the page and its observed queries; keyword repetition is not the deliverable.','',
'Build clear navigation between Lucknow’s hub, wedding, tempo, airport and relevant corridors. Prefer the correct final canonical URL. Check enquiries preserve the service and trip selected. Wedding work should separate guest transfers from bridal vehicle hire; group work should explain usable seats and luggage rather than publishing unsupported capacity claims.','',
'## 5. All city pages and all service families','',
table(['Family','Scope','Work required'],workstreams),'',
'For each of the 46 city hubs: record primary visitor task, current leading queries when available, actual service coverage, factual gaps, title/description decision, competing sibling pages, internal links, booking checks and LCP results. Mark a page “reviewed—keep” where its existing content already answers the task; do not add filler to satisfy an edit count.','',
'For every service family, test one representative shared component and all changed URLs, then broaden the check when a finding can affect siblings. Keep one-way, return, local package and activity-package pricing distinct. Preserve existing event names and the Week 51 deduplication/capture-form behaviour.','',
'## 6. Initial corridor shortlist','',
'This is a combined business-priority and historical-demand shortlist, not a claim that these are today’s top-converting routes. Fresh GSC and enquiry data can reorder it. Resolve aliases before editing.','',
table(['Requested corridor','Working canonical','Historical clicks','Position'],[[r['Requested path'],r['Working canonical'],r['Historical clicks'] if r['Historical clicks'] is not None else 'Unavailable',r['Historical position'] if r['Historical position'] is not None else 'Unavailable'] for r in routes]),'',
'Review pickup and final drop, departure constraints, trip type, realistic journey advice, actual vehicle choice and complete quote inclusions. Link to useful distance and toll resources without duplicating those entire pages. Protect the previous distance-page pilot until its results can be compared fairly.','',
'## 7. Exactly ten new blogs','',
'The proposed slugs do not exist in the local route list. That is only an exact-URL check: compare the full existing articles before final approval of each brief. If a brief duplicates an existing answer, narrow or replace it with a distinct traveller problem while retaining the total of ten new articles. Do not republish existing text under a different slug.','',
table(['ID','Proposed article','Focus'],[[b['ID'],b['Proposed title'],b['Intent']] for b in blogrows]),'']
for b in blogrows:md += ['### '+b['ID']+' — '+b['Proposed title'],'',b['Brief'],'','Proposed URL: '+b['Proposed URL']+'. Candidate main-site links: '+b['Candidate booking links']+'.','', 'Acceptance: the article solves a distinct traveller decision, uses verified changing facts, includes an original practical checklist or comparison, has accurate metadata, and links only where the reader benefits. Preserve relevant enquiry tracking. No arbitrary word count or unverified ranking promise.','']
md += ['## 8. LCP testing and fixes','',
'Begin with the actual LCP element and its timing. On all 46 city hubs, collect a mobile and desktop baseline. Use three comparable cold-load Lighthouse runs per device and retain the median plus run spread. Add the 13 corridor candidates and at least one representative URL per other service family; record their source URL, canonical and test environment. Test the Adventure homepage, an activity hub, an activity detail, a package and the booking page. Deduplicate overlapping URLs.','',
'Use the same browser/Lighthouse version, viewport, throttling, server mode and cache conditions for before/after comparisons. Test production builds, not development mode. Record TTFB, the LCP resource, resource-load delay, download time and render delay. Field data, where available, is a separate 28-day measurement; URL-level and origin-level results must be labelled correctly.','',
'The field objective is LCP at or below 2.5 seconds at the 75th percentile, assessed separately for mobile and desktop. A laboratory median below that threshold is useful engineering evidence but does not prove a field Core Web Vitals pass. Google’s metric guidance: https://web.dev/articles/lcp .','',
'Investigate hero image sizing, actual priority, image discovery, font loading, first-render visibility, avoidable hydration and server delays according to evidence. Do not blindly preload all images or remove tracking to improve a score. Preserve layout stability, meaningful image alternatives, keyboard use, enquiries and readable content. Document an unchanged or worse result honestly.','',
'After a shared fix, rerun the same suite and check all changed URLs for layout and rendering problems. Keep a separate daily measurement record for field trends; do not label a single fast run as the site’s final LCP.','',
'## 9. Adventure redesign brief','',
'Retain the Adventure identity and its distinction from cab booking. The local design uses warm paper, teal/rust accents, several fonts and animated hero content. Improve the browsing and booking decisions before adding visual effects. The proposed direction is a clearer mobile-first activity site with strong genuine photography, readable type, consistent activity cards and concise enquiry paths.','',
'Homepage: lead with location and what can be booked, followed by clear activity/category choices, a practical way to compare options, genuine supporting evidence and a simple next action. Keep essential heading and hero information visible on the first render. Retain useful existing sections, relocating them if needed rather than dropping content.','',
'Activity/category pages: show relevant options with actual duration, location, verified starting-price basis and inclusions. Detail pages should explain meeting point, what the experience includes, suitability restrictions confirmed by the operator, cancellation/weather conditions and the booking process. Never invent availability, safety credentials, reviews or superlatives. The local source contains “India’s highest bungee” and a 4.9/10,000-review claim: request supporting evidence before carrying those claims into the redesigned presentation.','',
'Enquiry journey: keep the selected activity, group size and date where the existing form supports them. Distinguish enquiry from confirmed booking. Keep Adventure’s own contact number and email; do not overwrite them with the cab business’s details. Test form validation and WhatsApp handoff without submitting fabricated leads to production.','',
'Measured performance candidates: Hero.tsx animates initial heading opacity from zero; SmartImage.tsx starts images at opacity zero until onLoad and fades them in; layout.tsx loads four font families. These are candidates for investigation, not established causes of poor LCP. Measure first, then ensure essential content does not wait on decorative animation.','',
'The local Adventure analytics configuration currently has empty GA4 ID, Search Console HTML verification token and lead-sheet endpoint. Verify live configuration and existing account ownership before treating tracking as active or broken. Search Console may already be verified through DNS; an empty HTML token alone does not establish lack of ownership. Use authorised real settings, and never copy IDs blindly.','',
'Redesign release checks: all existing public routes preserved; metadata/canonicals and sitemap remain correct; mobile and desktop navigation work; no unintended overflow; form state and contacts are correct; analytics configuration is verified; before/after LCP is recorded. The local manifest has 91 non-internal routes, including auxiliary output: individually distinguish pages from feeds and sitemaps in final QA.','',
'## 10. Relevant links between the two sites','',
'Link when a reader needs the other service: a traveller arranging a Rishikesh cab may need rafting or camping information; an activity visitor may need a pickup or return cab. These are related-brand navigation links, not independent editorial endorsements or an automatic authority uplift. Do not add a site-wide grid of exact-match reciprocal links or repeat links simply to inflate a count. Google’s relevant policy reference: https://developers.google.com/search/docs/essentials/spam-policies .','',
table(['Main-site candidate','Adventure candidate','Reader reason'],[[r['Main-site path'],r['Adventure path'],r['Reader reason']] for r in crossrows]),'',
'Both directions are in scope, but placement must match each page’s content. Confirm live status and final canonicals before insertion; candidate presence in a local build is not live verification. Use natural, descriptive anchors and avoid linking a generic activity to an unrelated city page. Keep existing relevant brand links.','',
'Record source page, destination URL, anchor, placement and purpose for every new link. Track cross-site navigation as an engagement action, not a booking conversion. If both properties share an authorised measurement setup, verify attribution and self-referrals; otherwise measure outbound and arrival behaviour separately without claiming a unified funnel.','',
'## 11. Weekly sequence','',
table(['Stage','Main site','Adventure / verification'],[['1','Fresh GSC, canonical map, city and family inventory, first blog overlap reviews','Inspect live release, freeze routes and contacts, capture design and LCP baselines'],['2','Lucknow city, wedding, tempo and route reviews; draft blogs B01–B03','Homepage and mobile navigation design; first-render visibility prototype'],['3','All-city reviews and selected route edits; draft B04–B05','Activity listing and detail templates; check actual content and claims'],['4','One-way, round-trip, local/outstation and group-vehicle work; draft B06–B07','Booking/enquiry redesign and form-state checks'],['5','Wedding and Char Dham review; draft B08–B10','Relevant cross-site links and tracking configuration checks'],['6','Finish individual review records; verify all ten blogs; rerun LCP after justified fixes','Responsive QA, route preservation and production-build checks'],['7','Controlled release of completed verified work; per-implementation MD and pending queue','Live smoke tests, tracking confirmation and post-release monitoring setup']]),'',
'These stages are a sequence, not a claim that staffing or review capacity is unlimited. If facts or access remain unavailable, carry the exact page or task forward with an owner and reason. Do not mark unverified pages complete merely because the week ends.','',
'## 12. Deliverables and acceptance','',
'Deliver ten distinct published articles once they pass editorial checks; a decision record for every city hub; an individual work queue for all service families; verified high-priority page improvements; an Adventure redesign with existing routes and useful content retained; contextual links in both directions; and comparable LCP evidence. Publish only work that passes its gates.','',
'Each implementation Markdown file must identify the problem, affected URLs, evidence window, exact sections/files changed, factual sources, screenshots where useful, test results, tracking preservation, deployment status and unresolved limitations. A shared template record must list every affected URL and distinguish individual checks from sampled checks.','',
'Measure organic clicks, relevant query/page rankings, CTR by comparable intent/position, qualified enquiry actions and field/lab performance separately. No numerical click target is included in the plan. Ranking uplift is an outcome to observe, not something to assert from completed edits.','',
'Files: week52_plan.md and week52_plan.docx; week52_work_register.xlsx; all-pages-work-register.csv; city-pages.csv; blog-publication-plan.csv; cross-site-link-plan.csv; LCP-test-register.csv; individual blog briefs; implementation-record-template.md. All are planning records until implementation evidence is added.','']
from week52_scope_additions import extend_plan
additional_sheets, additional_content = extend_plan(md, inventory, known, adventure, table)
(OUT/'adventure-seo-backlink-sightseeing-plan.md').write_text('\n'.join(additional_content),encoding='utf-8')
(OUT/'week52_plan.md').write_text('\n'.join(md),encoding='utf-8')
def csvout(name,rows):
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
csvout('all-pages-work-register.csv',inventory);csvout('city-pages.csv',cityrows);csvout('blog-publication-plan.csv',blogrows);csvout('cross-site-link-plan.csv',crossrows)
for _, rows, filename in additional_sheets: csvout(filename, rows)
lcprows=[]
mainTests=list(dict.fromkeys(citypaths+[r['Working canonical'] for r in routes]+['/wedding/lucknow','/chardham-yatra','/tempo-traveller/lucknow','/one-way-cab/delhi-to-dehradun','/round-trip/delhi-to-agra','/airport-service/lucknow','/blog/lucknow-nawabi-heritage-city-tour-guide-2026']))
mainTests=list(dict.fromkeys(mainTests+['/sightseeing','/sightseeing/lucknow','/sightseeing/lucknow-heritage-city-tour','/sightseeing/agra','/sightseeing/rishikesh']))
for site,paths in [('https://www.trivenicabs.in',mainTests),('https://adventure.trivenicabs.in',['/','/river-rafting-in-rishikesh','/camping-in-rishikesh','/bungee-rafting-camping-combo-rishikesh','/book-now','/rishikesh-haridwar-combo-tour'])]:
    for p in paths:
        for device in ['Mobile','Desktop']:lcprows.append({'URL':site+p,'Device':device,'Local existence':p in (known if site.endswith('www.trivenicabs.in') else adventure),'Phase':'Before and after, same conditions','Runs':3,'LCP median ms':None,'Run spread':None,'LCP element':None,'Field scope':'Pending URL/origin check','Field p75 ms':None,'Status':'Planned - not measured'})
csvout('LCP-test-register.csv',lcprows)
wb=Workbook();wb.remove(wb.active)
for name,rows in [('All pages',inventory),('City hubs',cityrows),('Priority routes',routes),('Ten blogs',blogrows),('Cross-site links',crossrows),('LCP tests',lcprows)]+[(name, rows) for name, rows, _ in additional_sheets]:
    ws=wb.create_sheet(name);h=list(rows[0]);ws.append(h)
    for r in rows:ws.append([r.get(k) for k in h])
    ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='16334D')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=45 if any(k in str(col[0].value).lower() for k in ['path','url','brief','action','title','link']) else 24
        for c in col[1:]:c.alignment=Alignment(wrap_text=True,vertical='top')
wb.save(OUT/'week52_work_register.xlsx')
briefdir=OUT/'blog-briefs';briefdir.mkdir(exist_ok=True)
for b in blogrows:(briefdir/(b['ID']+'.md')).write_text('# '+b['ID']+' — '+b['Proposed title']+'\n\n'+'\n\n'.join('**'+k+':** '+str(v) for k,v in b.items())+'\n\n## Publication evidence\n\n- Existing-content overlap review: pending\n- Author/editor: unassigned\n- Verified facts and sources: pending\n- Final title and URL: pending\n- Links and canonical checks: pending\n- Mobile and tracking checks: pending\n- Deployment: not published\n',encoding='utf-8')
(OUT/'implementation-record-template.md').write_text('# Week 52 implementation — title\n\n## Purpose and scope\n\n- Site and exact URLs:\n- Individual page/family IDs:\n- Baseline and evidence window:\n- Problem verified:\n\n## Changes\n\n- Exact sections and files changed:\n- Visitor benefit:\n- Factual sources and operations checks:\n- Existing content, routes and tracking preserved:\n\n## Verification\n\n- Build and relevant regression results:\n- Before/after LCP with environment and run spread:\n- Mobile/desktop screenshots:\n- Booking and enquiry state:\n- GA4 / Search Console / event deduplication checks:\n- Cross-site links and destination checks:\n\n## Release and follow-up\n\n- Implemented / verified / deployed status:\n- Live URLs and release evidence:\n- Remaining work and owner:\n- Later GSC and enquiry observations:\n\nDo not mark an unmeasured result as passed.\n',encoding='utf-8')
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.75)
doc.styles['Normal'].font.size=Pt(10.5)
lines='\n'.join(md).splitlines();i=0
while i<len(lines):
    line=lines[i]
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].startswith('|'):
            cells=[x.strip() for x in lines[i].strip('|').split('|')]
            if not all(x=='---' for x in cells):block.append(cells)
            i+=1
        t=doc.add_table(rows=1,cols=len(block[0]));t.style='Light Shading Accent 1'
        for j,c in enumerate(block[0]):t.rows[0].cells[j].text=c
        for row in block[1:]:
            cells=t.add_row().cells
            for j,c in enumerate(row):cells[j].text=c
        continue
    if line.startswith('# '):doc.add_heading(line[2:],0)
    elif line.startswith('## '):doc.add_heading(line[3:],1)
    elif line.startswith('### '):doc.add_heading(line[4:],2)
    elif line:doc.add_paragraph(line.replace('**',''))
    i+=1
sec.footer.paragraphs[0].text='TRIVENI | WEEK 52 | IMPLEMENTATION PLAN'
doc.save(OUT/'week52_plan.docx')
(OUT/'README.md').write_text('# Week 52\n\n[Detailed plan](week52_plan.md) · [Word](week52_plan.docx) · [PDF](week52_plan.pdf) · [Work register](week52_work_register.xlsx)\n\nScope: 10 blogs total; Lucknow; 46 city hubs; priority routes; wedding; Char Dham; group vehicles; local, outstation, one-way and round-trip cabs; LCP testing; Adventure redesign and useful links between both sites.\n\nAll records are planned, not completed work. Each implementation will receive its own Markdown evidence file using the included template. No production source or deployment was changed by this planning pass.\n',encoding='utf-8')
summary={'blogs':len(blogrows),'city_hubs':len(cityrows),'priority_corridors':len(routes),'adventure_manifest_routes':len(adventure),'combined_candidate_records':len(inventory),'lcp_url_device_cases':len(lcprows),'plan_words':len(' '.join(md).split()),'status':'Planning only; implementation and performance testing pending','source_window':'2026-06-24 to 2026-09-23','evidence_limits':'Local inventories and source review; no current Adventure GSC export or live performance measurements'}
(OUT/'planning-summary.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
print(json.dumps(summary,indent=2))
